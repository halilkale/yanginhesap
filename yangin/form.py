"""İşveren veri toplama formu (Excel) üretimi ve doldurulmuş formun programa aktarılması."""
from __future__ import annotations

import datetime as _dt
import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import alanlar as A
from . import etiketler as E
from . import tablolar as T
from .modeller import Girdi

LACIVERT = "1F3864"
SARI = "FFF2CC"
GRI = "F2F2F2"
KIM_RENK = {"Z": "F8CBAD", "V": "FFE699", "M": "BDD7EE"}
KIM_AD = {"Z": "İŞVEREN — zorunlu", "V": "İŞVEREN — varsa / biliyorsa", "M": "MÜHENDİS / montaj firması"}

_ince = Side(style="thin", color="A6A6A6")
KENAR = Border(left=_ince, right=_ince, top=_ince, bottom=_ince)
BASLIK_FILL = PatternFill("solid", fgColor=LACIVERT)
BASLIK_FONT = Font(bold=True, color="FFFFFF")

ANAHTAR_SUTUN = 10   # J: gizli program alan adı
HEADER_SATIR = 4
ILK_VERI_SATIRI = 5
TABLO_BASLIK = 4
TABLO_ILK_VERI = 7


# ---------------------------------------------------------------------------
# Şablon
# ---------------------------------------------------------------------------
class _Listeler:
    """Gizli 'Listeler' sayfasında seçenek listelerini tutar; doğrulama için aralık döndürür."""

    def __init__(self, wb: Workbook):
        self.ws = wb.create_sheet("Listeler")
        self.ws.sheet_state = "hidden"
        self.sutun = 0
        self.aralik: dict[str, str] = {}

    def ekle(self, ad: str, degerler: list[str]) -> str:
        if ad in self.aralik:
            return self.aralik[ad]
        self.sutun += 1
        for i, v in enumerate(degerler, 1):
            self.ws.cell(i, self.sutun, v)
        harf = get_column_letter(self.sutun)
        self.aralik[ad] = f"=Listeler!${harf}$1:${harf}${len(degerler)}"
        return self.aralik[ad]


def _dogrulama(ws, hucre_araligi: str, tip: str, liste: _Listeler, secenek=None, ad: str = ""):
    if tip == "evet":
        dv = DataValidation(type="list", formula1='"Evet,Hayır"', allow_blank=True)
    elif tip == "secim":
        dv = DataValidation(type="list", formula1=liste.ekle(ad, list(secenek.values())), allow_blank=True)
    elif tip == "tesis":
        dv = DataValidation(type="list", formula1=liste.ekle("tesis", A.TESIS_ETIKET), allow_blank=True)
    elif tip == "tam":
        dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    elif tip == "sayi":
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    else:
        return
    if tip in ("tam", "sayi"):
        dv.errorTitle, dv.error = "Geçersiz değer", "Lütfen 0 veya daha büyük bir sayı yazın."
    else:
        dv.errorTitle, dv.error = "Geçersiz seçim", "Lütfen açılır listeden seçin."
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(hucre_araligi)


def _sayfa_basligi(ws, baslik: str, giris: str, genislik_sutun: int):
    ws["A1"] = baslik
    ws["A1"].font = Font(bold=True, size=14, color=LACIVERT)
    ws["A2"] = giris
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=genislik_sutun)
    ws.row_dimensions[2].height = 42


def _talimat(wb: Workbook):
    ws = wb.active
    ws.title = "Talimat"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 110
    ws["B1"] = "YANGIN KORUMA HESABI — İŞVEREN VERİ TOPLAMA FORMU"
    ws["B1"].font = Font(bold=True, size=16, color=LACIVERT)
    ws["B2"] = "Binaların Yangından Korunması Hakkında Yönetmelik kapsamında yapılacak hesaplar için işletmeden alınması gereken bilgiler"
    ws["B2"].font = Font(italic=True, color="595959")
    satirlar = [
        ("Bu form ne için?", "Mühendisin yangın korunum hesaplarını (kaçış genişlikleri, kompartıman, sprinkler, yangın suyu deposu, dolap/hidrant vb.) yapabilmesi için işletmenin bilmesi ve belgelerden çıkarabileceği bilgileri toplar. "
                             "Doldurulan form 'Yangın Koruma Hesap' programına tek tıkla aktarılır; bilgiler yeniden yazılmaz."),
        ("Nasıl doldurulur?", "1) Sayfaları soldan sağa doldurun (1_Genel … 7_Yapi_Elemanlari, sonra T1 … T6 tabloları). \n"
                              "2) Yalnızca SARI hücrelere yazın. Gri/sabit hücrelere dokunmayın, sütun veya satır silmeyin, sayfa adlarını değiştirmeyin.\n"
                              "3) Açılır listesi olan hücrelerde listeden seçin. Sayı isteyen hücrelere yalnızca sayı yazın (birim yazmayın; ondalık için virgül veya nokta).\n"
                              "4) Bilmediğiniz veya sizde olmayan bir bilgiyi BOŞ bırakın; uydurmayın. Boş bırakılan bilgi mühendise 'eksik veri' olarak bildirilir.\n"
                              "5) Bilginin nereden bulunacağı 'Belge / kaynak' sütununda yazılıdır. Belgelerin kopyalarını 'Belgeler' sayfasındaki listeye göre ekleyin."),
        ("Kim dolduracak?", "'Kim doldurur' sütunu: İŞVEREN-zorunlu = mutlaka sizden bekleniyor; İŞVEREN-varsa = sizde varsa/biliyorsanız; "
                            "MÜHENDİS = proje müellifi, montaj-bakım firması veya mühendis doldurur; işveren boş bırakabilir veya ilgili firmaya yönlendirebilir."),
        ("Renkler", "SARI hücre: doldurun.   TURUNCU 'Kim' etiketi: sizden bekleniyor.   MAVİ 'Kim' etiketi: teknik firma/mühendis.   GRİ satır: örnek, silmeyin (örnek satırlar programa aktarılmaz)."),
        ("Ölçü ve birimler", "Uzunluk metre (m), alan metrekare (m²), süre dakika (dk), debi litre/dakika veya m³/saat, basınç/yükseklik mSS (≈ metre su sütunu). Birimler her sorunun yanında yazılıdır."),
        ("Tablolar (T1–T6)", "Her satır bir kat, bir bölüm, bir merdiven, bir kapı vb. Gerektiği kadar satır doldurun; kullanmadığınız satırları boş bırakın. 'Örnek' yazan satırı silmeyin. "
                              "T1 ve T2'de KAT ADLARI aynı yazılmalıdır (Zemin kat, 1. kat ...)."),
        ("Gizlilik", "Form yalnızca yangın güvenliği hesapları için kullanılır. Ticari sır niteliğindeki üretim ayrıntılarını yazmanız gerekmez; faaliyet türü ve kullanılan madde sınıfları yeterlidir."),
        ("Sonrası", "Formu ve belge kopyalarını mühendise iletin. Mühendis formu programa aktarır, eksik/uygunsuz hususları rapor eder; sonuçlar Excel ve Word raporu olarak verilir."),
    ]
    r = 4
    for b, m in satirlar:
        ws.cell(r, 2, b).font = Font(bold=True)
        c = ws.cell(r, 3, m)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 2).alignment = Alignment(vertical="top")
        ws.row_dimensions[r].height = max(30, 15 * (len(m) // 105 + 1 + m.count("\n")))
        r += 1
    r += 1
    ws.cell(r, 2, "Sayfalar").font = Font(bold=True, color=LACIVERT)
    r += 1
    for s in A.SAYFALAR:
        ws.cell(r, 2, s.ad)
        ws.cell(r, 3, s.baslik + " — " + s.giris)
        r += 1
    for t in A.TABLO_SAYFALARI:
        ws.cell(r, 2, t.ad)
        ws.cell(r, 3, t.baslik)
        r += 1
    ws.cell(r, 2, "Belgeler")
    ws.cell(r, 3, "İşletmeden istenecek belge kontrol listesi")
    ws.cell(r + 1, 2, "Notlar")
    ws.cell(r + 1, 3, "Forma sığmayan ek bilgiler")


def _alan_sayfasi(wb: Workbook, s: A.Sayfa, liste: _Listeler):
    ws = wb.create_sheet(s.ad)
    _sayfa_basligi(ws, s.baslik, s.giris, 8)
    basliklar = ["No", "Soru / bilgi", "Nasıl doldurulur", "Birim", "CEVABINIZ", "Örnek", "Kim doldurur", "Belge / kaynak"]
    for j, b in enumerate(basliklar, 1):
        c = ws.cell(HEADER_SATIR, j, b)
        c.fill, c.font, c.border = BASLIK_FILL, BASLIK_FONT, KENAR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(HEADER_SATIR, ANAHTAR_SUTUN, "alan")
    for i, a in enumerate(s.alanlar):
        r = ILK_VERI_SATIRI + i
        degerler = [i + 1, a.baslik, a.aciklama, a.birim, None, a.ornek, KIM_AD[a.kim], a.kaynak]
        for j, v in enumerate(degerler, 1):
            c = ws.cell(r, j, v)
            c.border = KENAR
            c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if j in (1, 4) else "left")
        cevap = ws.cell(r, 5)
        cevap.fill = PatternFill("solid", fgColor=SARI)
        cevap.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 6).font = Font(italic=True, color="7F7F7F")
        ws.cell(r, 7).fill = PatternFill("solid", fgColor=KIM_RENK[a.kim])
        ws.cell(r, ANAHTAR_SUTUN, a.alan or "")
        _dogrulama(ws, f"E{r}", a.tip, liste, a.secenek, ad=a.alan or "")
        uzun = max(len(a.baslik) // 38, len(a.aciklama) // 52, len(a.kaynak) // 28, 1)
        ws.row_dimensions[r].height = min(15 * (uzun + 1), 120)
    for j, w in enumerate([5, 44, 62, 9, 30, 22, 20, 30], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.column_dimensions[get_column_letter(ANAHTAR_SUTUN)].hidden = True
    ws.freeze_panes = ws.cell(ILK_VERI_SATIRI, 3)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _tablo_sayfasi(wb: Workbook, t: A.TabloSayfa, liste: _Listeler):
    ws = wb.create_sheet(t.ad)
    kolonlar = E.TABLOLAR[t.anahtar]
    _sayfa_basligi(ws, t.baslik, t.giris + f"   [Kim: {KIM_AD[t.kim]}.  Kaynak: {t.kaynak}]", max(len(kolonlar), 6))
    ws.cell(3, 1, "Satır 6 örnektir (programa aktarılmaz). Verilerinizi 7. satırdan itibaren yazın.").font = Font(italic=True, color="C00000")
    for j, (anahtar, baslik, tur, secenek, genislik) in enumerate(kolonlar, 1):
        c = ws.cell(TABLO_BASLIK, j, baslik)
        c.fill, c.font, c.border = BASLIK_FILL, BASLIK_FONT, KENAR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        a = ws.cell(5, j, t.aciklamalar.get(anahtar, ""))
        a.font = Font(italic=True, size=8, color="595959")
        a.alignment = Alignment(wrap_text=True, vertical="top")
        a.border = KENAR
        a.fill = PatternFill("solid", fgColor=GRI)
        ornek = t.ornek.get(anahtar, "")
        if j == 1:
            ornek = f"ÖRN: {ornek}"
        e = ws.cell(6, j, ornek)
        e.font = Font(italic=True, color="7F7F7F")
        e.fill = PatternFill("solid", fgColor=GRI)
        e.border = KENAR
        ws.column_dimensions[get_column_letter(j)].width = max(14, min(genislik / 6.5, 34))
        ws.cell(TABLO_BASLIK, j).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        aralik = f"{get_column_letter(j)}{TABLO_ILK_VERI}:{get_column_letter(j)}{TABLO_ILK_VERI + t.satir - 1}"
        if tur == "b":
            _dogrulama(ws, aralik, "evet", liste)
        elif tur == "c":
            _dogrulama(ws, aralik, "secim", liste, {s: s for s in secenek}, ad=f"{t.anahtar}_{anahtar}")
        elif tur in ("f", "i"):
            _dogrulama(ws, aralik, "sayi" if tur == "f" else "tam", liste)
        for r in range(TABLO_ILK_VERI, TABLO_ILK_VERI + t.satir):
            h = ws.cell(r, j)
            h.fill = PatternFill("solid", fgColor=SARI)
            h.border = KENAR
    ws.row_dimensions[4].height = 34
    ws.row_dimensions[5].height = 118
    ws.freeze_panes = ws.cell(TABLO_ILK_VERI, 1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _belgeler(wb: Workbook, liste: _Listeler):
    ws = wb.create_sheet("Belgeler")
    _sayfa_basligi(ws, "Belge kontrol listesi", "Aşağıdaki belgelerin kopyalarını forma ekleyin. 'Durum' sütununu işaretleyin.", 6)
    for j, b in enumerate(["No", "Belge", "Neden isteniyor", "Kimden alınır", "Zorunluluk", "Durum", "Not"], 1):
        c = ws.cell(HEADER_SATIR, j, b)
        c.fill, c.font, c.border = BASLIK_FILL, BASLIK_FONT, KENAR
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    dv = DataValidation(type="list", formula1='"Teslim edildi,Yok,Bu tesiste geçerli değil,Sonra iletilecek"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, (belge, neden, kimden, zor) in enumerate(A.BELGELER):
        r = ILK_VERI_SATIRI + i
        for j, v in enumerate([i + 1, belge, neden, kimden, zor, None, None], 1):
            c = ws.cell(r, j, v)
            c.border = KENAR
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 6).fill = PatternFill("solid", fgColor=SARI)
        ws.cell(r, 7).fill = PatternFill("solid", fgColor=SARI)
        ws.cell(r, 5).fill = PatternFill("solid", fgColor=KIM_RENK["Z"] if zor == "Zorunlu" else KIM_RENK["V"])
        dv.add(f"F{r}")
        ws.row_dimensions[r].height = 46
    for j, w in enumerate([5, 58, 52, 28, 12, 22, 30], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A5"


def _notlar(wb: Workbook):
    ws = wb.create_sheet("Notlar")
    _sayfa_basligi(ws, "Ek bilgiler ve notlar", "Forma sığmayan veya mühendisin bilmesi gereken diğer hususları (özel üretim süreçleri, planlanan genişleme, bilinen yangın olayları, itfaiye uyarıları vb.) yazın.", 3)
    ws.column_dimensions["A"].width = 120
    for r in range(4, 24):
        c = ws.cell(r, 1)
        c.fill = PatternFill("solid", fgColor=SARI)
        c.border = KENAR
        c.alignment = Alignment(wrap_text=True, vertical="top")


def sablon_olustur() -> bytes:
    wb = Workbook()
    _talimat(wb)
    liste = _Listeler(wb)
    for s in A.SAYFALAR:
        _alan_sayfasi(wb, s, liste)
    for t in A.TABLO_SAYFALARI:
        _tablo_sayfasi(wb, t, liste)
    _belgeler(wb, liste)
    _notlar(wb)
    wb.move_sheet("Listeler", offset=len(wb.sheetnames))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ---------------------------------------------------------------------------
# İçe aktarma
# ---------------------------------------------------------------------------
def _sayi(v):
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", ".")
    return float(s)


def _evet(v):
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, bool):
        return v
    t = str(v).strip().lower()
    if t in ("evet", "e", "var", "true", "1", "yes"):
        return True
    if t in ("hayır", "hayir", "h", "yok", "false", "0", "no"):
        return False
    raise ValueError(v)


def formdan_oku(veri: bytes) -> tuple[Girdi, list[str], list[str]]:
    """Doldurulmuş formu okur. (girdi, uyarılar, eksik_zorunlu) döndürür."""
    wb = load_workbook(io.BytesIO(veri), data_only=True)
    g = Girdi(tarih=_dt.date.today().strftime("%d.%m.%Y"))
    uyari: list[str] = []
    eksik: list[str] = []
    tesis_secimi = None
    tehlike_acik = False
    for s in A.SAYFALAR:
        if s.ad not in wb.sheetnames:
            uyari.append(f"'{s.ad}' sayfası bulunamadı (sayfa adı değiştirilmiş olabilir).")
            continue
        ws = wb[s.ad]
        anahtar_satir = {}
        for r in range(ILK_VERI_SATIRI, ILK_VERI_SATIRI + len(s.alanlar) + 5):
            k = ws.cell(r, ANAHTAR_SUTUN).value
            if k:
                anahtar_satir[k] = r
        for a in s.alanlar:
            if not a.alan:
                continue
            r = anahtar_satir.get(a.alan if a.alan != "@tesis" else "@tesis")
            if r is None:
                continue
            ham = ws.cell(r, 5).value
            bos = ham is None or str(ham).strip() == ""
            if bos:
                if a.kim == "Z" and a.tip != "bilgi":
                    eksik.append(f"{s.ad}: {a.baslik}")
                continue
            try:
                if a.tip == "tesis":
                    tesis_secimi = A.TESIS_ETIKET.index(str(ham).strip())
                    continue
                if a.tip == "evet":
                    setattr(g, a.alan, _evet(ham))
                elif a.tip == "secim":
                    ters = {v: k for k, v in a.secenek.items()}
                    etiket = str(ham).strip()
                    if etiket not in ters:
                        raise ValueError(etiket)
                    setattr(g, a.alan, ters[etiket])
                    if a.alan == "tehlike":
                        tehlike_acik = True
                elif a.tip == "sayi":
                    setattr(g, a.alan, _sayi(ham))
                elif a.tip == "tam":
                    setattr(g, a.alan, int(round(_sayi(ham))))
                else:
                    setattr(g, a.alan, str(ham).strip())
            except (ValueError, TypeError):
                uyari.append(f"{s.ad}: '{a.baslik}' alanındaki değer anlaşılamadı ({ham!r}); varsayılan değer kullanıldı.")
    if tesis_secimi is not None and not tehlike_acik:
        g.tehlike = T.EK1_TESISLER[tesis_secimi][2]
    # --- tablolar
    satirlar: dict[str, list[dict]] = {}
    for t in A.TABLO_SAYFALARI:
        liste: list[dict] = []
        if t.ad not in wb.sheetnames:
            uyari.append(f"'{t.ad}' sayfası bulunamadı.")
            satirlar[t.anahtar] = liste
            continue
        ws = wb[t.ad]
        kolonlar = E.TABLOLAR[t.anahtar]
        for r in range(TABLO_ILK_VERI, ws.max_row + 1):
            hucreler = [ws.cell(r, j).value for j in range(1, len(kolonlar) + 1)]
            if all(h is None or str(h).strip() == "" for h in hucreler):
                continue
            if str(hucreler[0] or "").upper().startswith("ÖRN"):
                continue
            d = {}
            for (anahtar, baslik, tur, secenek, _g), h in zip(kolonlar, hucreler):
                bos = h is None or str(h).strip() == ""
                if tur == "b":
                    try:
                        v = _evet(h)
                    except ValueError:
                        uyari.append(f"{t.ad} satır {r}: '{baslik}' için Evet/Hayır bekleniyor ({h!r}).")
                        v = None
                    d[anahtar] = v if v is not None else (anahtar in ("sayilir", "yone"))
                elif tur in ("f", "i"):
                    try:
                        d[anahtar] = _sayi(h) or 0
                    except ValueError:
                        uyari.append(f"{t.ad} satır {r}: '{baslik}' sayı olmalı ({h!r}).")
                        d[anahtar] = 0
                elif tur == "c":
                    d[anahtar] = str(h).strip() if not bos else secenek[0]
                    if not bos and str(h).strip() not in secenek:
                        uyari.append(f"{t.ad} satır {r}: '{baslik}' listede yok ({h!r}).")
                else:
                    d[anahtar] = "" if bos else str(h).strip()
            liste.append(d)
        satirlar[t.anahtar] = liste
    g = E.tablolari_uygula(g, satirlar["mahal"], satirlar["kat"], satirlar["merdiven"], satirlar["kapi"], satirlar["duman"], satirlar["tank"])
    mahal_katlari = {r["kat"] for r in satirlar["mahal"]}
    kat_katlari = {r["kat"] for r in satirlar["kat"]}
    for k in sorted(mahal_katlari - kat_katlari):
        uyari.append(f"T1_Mahaller'deki '{k}' katı T2_Katlar tablosunda yok; çıkış verisi girilmedi (kat adlarını aynı yazın).")
    if not satirlar["mahal"]:
        eksik.append("T1_Mahaller: hiç mahal girilmemiş (kullanıcı yükü hesaplanamaz)")
    return g, uyari, eksik


# ---------------------------------------------------------------------------
# Girdi -> doldurulmuş form (örnek form ve proje dışa aktarımı)
# ---------------------------------------------------------------------------
def doldurulmus_form(g: Girdi) -> bytes:
    """Verilen projeyle doldurulmuş form üretir (örnek form; programdan formu geri üretmek için)."""
    wb = load_workbook(io.BytesIO(sablon_olustur()))
    for s in A.SAYFALAR:
        ws = wb[s.ad]
        for i, a in enumerate(s.alanlar):
            r = ILK_VERI_SATIRI + i
            if not a.alan or a.alan == "@tesis":
                continue
            v = getattr(g, a.alan)
            if a.tip == "evet":
                v = "Evet" if v else "Hayır"
            elif a.tip == "secim":
                v = a.secenek[v]
            elif a.tip == "sayi" and isinstance(v, (int, float)) and float(v) == 0 and a.kim != "Z":
                v = None
            ws.cell(r, 5, v)
    # tablolar
    from .etiketler import CIKIS_TURLERI, YONLER
    veri = {
        "mahal": [{"kat": k.ad, "mahal": m.ad,
                   "tur": next((x for x, kat, _ in T.EK5A if abs(kat - m.katsayi) < 1e-9 and "Diğer" not in x), "Diğer / özel katsayı (elle gir)"),
                   "alan": m.alan, "ozel": m.katsayi if not any(abs(kat - m.katsayi) < 1e-9 for _, kat, _ in T.EK5A) else None,
                   "belirli": m.kisi_belirli, "sayilir": m.sayilir} for k in g.katlar for m in k.mahaller],
        "kat": [{"kat": k.ad, "cikis": k.cikis_sayisi, "tur": CIKIS_TURLERI[k.cikis_turu], "genislik": k.mevcut_genislik_cm, "tekil": k.tekil_cikis_genislik_cm,
                 "yon": YONLER[k.yon], "mesafe": k.en_uzak_mesafe_m, "kus": k.kus_ucusu_mesafe_m or None, "cikmaz": k.cikmaz_mesafe_m or None,
                 "diyagonal": k.mekan_diyagonal_m or None, "arasi": k.cikislar_arasi_mesafe_m or None} for k in g.katlar],
        "merdiven": [{"ad": m.ad, "gen": m.genislik_cm, "riht": m.rihts_mm, "basis": m.basamak_genislik_mm, "sbasamak": m.sahanlik_arasi_basamak, "skot": m.sahanlik_arasi_kot_cm,
                      "bas": m.bas_yuksekligi_cm, "kat": m.hizmet_verilen_kat, "kapi": m.kapi_dayanim_dk, "duvar": m.duvar_dayanim_dk, "kullanici": m.kullanici_sayisi_kat,
                      "dengeli": m.dengelenmis} for m in g.merdivenler],
        "kapi": [{"ad": k.ad, "gen": k.temiz_genislik_cm, "yuk": k.yukseklik_cm, "kanat": k.kanat_sayisi, "kisi": k.kisi_yuku, "yone": k.kacis_yonune_aciliyor,
                  "esik": k.esik_var, "kuvvet": k.acma_kuvveti_N or None} for k in g.kapilar],
        "duman": [{"ad": m.ad, "tur": m.tur, "alan": m.alan_m2, "yuk": m.yukseklik_m} for m in g.duman_mahalleri],
        "tank": [{"ad": t.ad, "sinif": t.sinif, "hacim": t.hacim_L, "tur": t.tur} for t in g.sivi_tanklar],
    }
    for t in A.TABLO_SAYFALARI:
        ws = wb[t.ad]
        for i, d in enumerate(veri[t.anahtar]):
            for j, (anahtar, _b, tur, _s, _g) in enumerate(E.TABLOLAR[t.anahtar], 1):
                v = d.get(anahtar)
                if tur == "b":
                    v = "Evet" if v else "Hayır"
                ws.cell(TABLO_ILK_VERI + i, j, v)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
