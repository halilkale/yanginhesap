"""Ayrıntılı kullanım kılavuzu (Word) üretimi."""
from __future__ import annotations

import datetime as _dt
import io
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import alanlar as A
from . import etiketler as E
from . import hesaplar as H
from . import tablolar as T
from .modeller import GEREKLI, KONTROL, KOSULLU, UYGUN, UYGUN_DEGIL, GEREKMEZ, BILGI
from .ornek import ornek_fabrika

LACIVERT = RGBColor(0x1F, 0x38, 0x64)
SURUM = "1.0"
RELEASE = "https://github.com/halilkale/yanginhesap/releases/tag/exe-son"
DEPO = "https://github.com/halilkale/yanginhesap"

DURUM_ACIKLAMA = [
    (UYGUN, "C6EFCE", "Girdiğiniz değer yönetmeliğin aradığı sınırı sağlıyor.", "Eylem gerekmez; rapora kanıt olarak girer."),
    (UYGUN_DEGIL, "FFC7CE", "Girdiğiniz değer yönetmelik sınırını sağlamıyor veya yasak bir durum var.", "Tasarım/uygulama düzeltilmeli (genişlik artırma, çıkış ekleme, sprinkler, kompartıman ayırma vb.). Açıklama sütunu nedenini ve gereken değeri yazar."),
    (GEREKLI, "DDEBF7", "Yönetmelik bu tesis için bir önlemi zorunlu kılıyor (ör. acil aydınlatma, itfaiye su alma ağzı, ekiplerin kurulması).", "Önlemin projede/işletmede var olduğunu belgeleyin. Önlem sizde zaten varsa mühendis 'UYGUN' olarak teyit eder."),
    (KOSULLU, "FFEB9C", "Sonuç bir şarta bağlı (ör. kompartıman alanı 'uygun yangın kontrol sistemleri' bulunduğu için sınırsız).", "Şartın sağlandığı, gerekirse itfaiye görüşüyle doğrulanmalı."),
    (KONTROL, "EDEDED", "Kontrolü yapılabilmesi için ölçülen/projedeki değer girilmemiş. Gerekli değer yine de hesaplanır ve açıklamada yazılır.", "İşverenden/müelliften eksik veriyi isteyin, forma girin, yeniden hesaplayın."),
    (GEREKMEZ, "F2F2F2", "Bu tesis için yönetmelik o önlemi zorunlu kılmıyor.", "Eylem gerekmez (isteğe bağlı iyileştirme olarak düşünülebilir)."),
    (BILGI, "FFFFFF", "Ara hesap, tablo değeri veya açıklayıcı bilgi.", "Rapora hesabın izlenebilmesi için girer."),
]


# ---------------------------------------------------------------------------
# Biçim yardımcıları
# ---------------------------------------------------------------------------
def _golge(cell, renk: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), renk)
    for eski in tcPr.findall(qn("w:shd")):
        tcPr.remove(eski)
    tcPr.insert_element_before(shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")


def _hucre(cell, metin, kalin=False, boyut=9, renk=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(metin))
    r.font.size = Pt(boyut)
    r.bold = kalin
    if renk:
        r.font.color.rgb = renk


def _tablo(doc, basliklar, satirlar, genislikler=None, boyut=8.5, baslik_renk="1F3864"):
    t = doc.add_table(rows=1, cols=len(basliklar))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, b in enumerate(basliklar):
        _hucre(t.rows[0].cells[j], b, kalin=True, boyut=boyut, renk=RGBColor(0xFF, 0xFF, 0xFF))
        _golge(t.rows[0].cells[j], baslik_renk)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trPr.append(h)
    for satir in satirlar:
        c = t.add_row().cells
        for j, v in enumerate(satir):
            _hucre(c[j], v, boyut=boyut)
    if genislikler:
        for row in t.rows:
            for j, w in enumerate(genislikler):
                row.cells[j].width = Cm(w)
    doc.add_paragraph()
    return t


def _p(doc, metin, kalin=False, italik=False, boyut=None, hizala=None, bosluk=6):
    p = doc.add_paragraph()
    r = p.add_run(metin)
    r.bold, r.italic = kalin, italik
    if boyut:
        r.font.size = Pt(boyut)
    if hizala:
        p.alignment = hizala
    p.paragraph_format.space_after = Pt(bosluk)
    return p


def _madde(doc, metin, kalin_baslangic: str = "", seviye=0):
    p = doc.add_paragraph(style="List Bullet" if seviye == 0 else "List Bullet 2")
    if kalin_baslangic:
        r = p.add_run(kalin_baslangic)
        r.bold = True
    p.add_run(metin)
    p.paragraph_format.space_after = Pt(2)
    return p


def _numara(doc, metin, kalin_baslangic=""):
    p = doc.add_paragraph(style="List Number")
    if kalin_baslangic:
        p.add_run(kalin_baslangic).bold = True
    p.add_run(metin)
    p.paragraph_format.space_after = Pt(2)
    return p


def _not(doc, metin, baslik="Not", renk="FFF2CC"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    c.text = ""
    p = c.paragraphs[0]
    p.add_run(f"{baslik}: ").bold = True
    p.add_run(metin)
    for r in p.runs:
        r.font.size = Pt(9.5)
    _golge(c, renk)
    doc.add_paragraph()


def _resim(doc, yol, aciklama, genislik=16.5):
    if yol and os.path.exists(yol):
        doc.add_picture(yol, width=Cm(genislik))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _p(doc, aciklama, italik=True, boyut=8.5, hizala=WD_ALIGN_PARAGRAPH.CENTER)


def _alan_numarasi(p, komut):
    r = p.add_run()
    for tip, metin in (("begin", None), (None, komut), ("end", None)):
        if tip:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), tip)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = metin
        r._r.append(e)


def _sayfa_sonu(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _fmt(v):
    if isinstance(v, float):
        return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".").rstrip("0").rstrip(",") if v != int(v) else f"{int(v):,}".replace(",", ".")
    return str(v)


def _bul(sonuclar, kalem, bolum=""):
    for r in sonuclar:
        if kalem in r.kalem and bolum in r.bolum:
            return r
    return None


# ---------------------------------------------------------------------------
# Ana üretici
# ---------------------------------------------------------------------------
def kilavuz_olustur(ekran_dizini: str | None = None) -> bytes:
    ekran = lambda ad: os.path.join(ekran_dizini, ad) if ekran_dizini else None
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    n = doc.styles["Normal"]
    n.font.name = "Calibri"
    n.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    n.font.size = Pt(10.5)
    for ad, boyut in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        st = doc.styles[ad]
        st.font.name = "Calibri"
        st.font.size = Pt(boyut)
        st.font.color.rgb = LACIVERT
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    # alt bilgi: sayfa numarası
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Yangın Koruma Hesap Uygulaması — Kullanım Kılavuzu   |   Sayfa ").font.size = Pt(8)
    _alan_numarasi(fp, "PAGE")

    # ---------------- Kapak
    for _ in range(5):
        doc.add_paragraph()
    _p(doc, "YANGIN KORUMA HESAP UYGULAMASI", kalin=True, boyut=26, hizala=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, "KULLANIM KILAVUZU", kalin=True, boyut=20, hizala=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    _p(doc, "Binaların Yangından Korunması Hakkında Yönetmelik kapsamında\nfabrika / endüstriyel yapılar için hesap ve kontrol programı",
       italik=True, boyut=13, hizala=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(4):
        doc.add_paragraph()
    _p(doc, f"Sürüm {SURUM}   —   {_dt.date.today().strftime('%d.%m.%Y')}", hizala=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, "Dayanak: Binaların Yangından Korunması Hakkında Yönetmelik (RG 19.12.2007/26735; 20.11.2021 tarihli 4825 sayılı Cumhurbaşkanı Kararı ile değişik) ve "
            "Binaların Yangından Korunması Hakkında Yönetmelik Kılavuzu (Çevre, Şehircilik ve İklim Değişikliği Bakanlığı, Aralık 2024)",
       boyut=9, hizala=WD_ALIGN_PARAGRAPH.CENTER)
    _sayfa_sonu(doc)

    # ---------------- İçindekiler (statik)
    doc.add_heading("İçindekiler", 1)
    for satir in ("1. Giriş: programın amacı ve kapsamı", "2. Çalışma akışı: baştan sona 8 adım", "3. Kurulum", "4. İşverenden veri toplama (Excel formu)",
                  "5. Program ekranı: düğmeler ve sekmeler", "6. Verilerin girilmesi: sekme sekme, alan alan", "7. Hesaplama ve sonuçların okunması",
                  "8. Raporlar (Excel ve Word) ve proje dosyası", "9. Hesap yöntemleri ve formüller", "10. Çözümlü örnek: Örnek Mobilya Fabrikası",
                  "11. Sorun giderme ve sık sorulan sorular", "12. Sınırlar, varsayımlar ve sorumluluk", "Ek-A. Terimler sözlüğü", "Ek-B. Yönetmelik maddesi — program çıktısı eşlemesi",
                  "Ek-C. Rapor teslim öncesi kontrol listesi"):
        _p(doc, satir, bosluk=2)
    _sayfa_sonu(doc)

    # ---------------- 1 Giriş
    doc.add_heading("1. Giriş: programın amacı ve kapsamı", 1)
    _p(doc, "Bu program, Binaların Yangından Korunması Hakkında Yönetmeliğin sayısal kriterlerini ve zorunluluk koşullarını, özellikle fabrika, imalathane ve depo gibi "
            "endüstriyel yapılar için, tek bir veri girişinden hesaplar. Sonuçları madde ve ek numarasıyla birlikte Excel ve Word raporu olarak verir.")
    doc.add_heading("1.1 Kimler için?", 2)
    _madde(doc, "Yangın güvenliği, mekanik/elektrik/inşaat mühendisleri ve mimarlar (proje kontrolü, ön boyutlandırma).")
    _madde(doc, "İş sağlığı ve güvenliği (İSG) uzmanları ve yangın güvenliği danışmanları (mevcut işletmenin yönetmeliğe uygunluğunun incelenmesi).")
    _madde(doc, "İşverenler / tesis yöneticileri (hangi bilgilerin gerektiğini görmek ve sonuçları takip etmek).")
    doc.add_heading("1.2 Program ne yapar?", 2)
    _tablo(doc, ["Konu", "Yapılan hesap / kontrol", "Yönetmelik dayanağı"], [
        ["Sınıflandırma", "Kullanım sınıfı, tehlike sınıfı (Ek-1 listesinden öneri), yüksek bina, tahliye projesi gerekliliği", "Madde 4, 7, 8, 14, 17-19, Ek-1"],
        ["Yangın dayanımı", "Taşıyıcı sistem ve kompartıman duvar/döşeme dayanım süreleri; çelik yalıtımı; beton paspayı; ahşap kalan kesit; cephe ve çatı", "Madde 23-28, Ek-3/B, Ek-3/C"],
        ["Kompartıman", "En fazla kompartıman alanı ve dipnot koşulları (kontrol sistemi, tek katlı)", "Madde 24, Ek-4"],
        ["Kaçış", "Kullanıcı yükü, çıkış sayısı, çıkış genişliği, kaçış mesafesi, çıkmaz koridor, çıkışlar arası mesafe, merdiven ve kapı ölçüleri", "Madde 31-47, 52, Ek-5/A-B, Ek-14"],
        ["Yağmurlama (sprinkler)", "Zorunluluk, tasarım debisi, yaklaşık başlık adedi", "Madde 96, Ek-8/B"],
        ["Yangın suyu", "Su deposu hacmi, toplam debi, pompa debisi/basma yüksekliği, pompa kontrolleri", "Madde 92-93, Ek-8/A-B-C"],
        ["Dolap, hidrant, söndürücü", "Zorunluluk, yaklaşık adetler, itfaiye su alma/verme bağlantısı", "Madde 94-97, 99"],
        ["Algılama ve aydınlatma", "Algılama, buton, acil aydınlatma süre/seviye, yönlendirme levhası", "Madde 70-77, 81, Ek-7"],
        ["Duman ve basınçlandırma", "Mekanik duman tahliyesi (hava debisi), merdiven basınçlandırma zorunluluğu ve asgari debi", "Madde 85-89"],
        ["Kazan, yakıt, trafo, jeneratör", "Kapı sayısı, yakıt depolama limitleri, havuzlama, oda dayanımı", "Madde 54-56, 65-66"],
        ["Tehlikeli maddeler", "LPG emniyet uzaklıkları, yanıcı sıvı eşdeğerlik/bildirim/izin/mesafe", "Madde 106-122, Ek-9…12"],
        ["Çevre ve erişim", "İtfaiye erişimi, orman alanı dış yangın bölgesi", "Madde 21-22"],
        ["Organizasyon", "Acil durum ekipleri ve asgari personel, tatbikat", "Madde 124-129"],
    ], [3.6, 9.2, 4.2])
    doc.add_heading("1.3 Program ne yapmaz?", 2)
    _madde(doc, "Boru çapı ve ayrıntılı hidrolik hesabı, yapısal yangın analizi (statik hesap), duman tahliye hacim hesabı yapmaz; yönetmelik bunlar için sayı vermez, standartlara atıf yapar. Bu konular ilgili satırlarda 'standarda göre tasarım' olarak belirtilir.")
    _madde(doc, "Projenin onayı, itfaiye görüşü veya yetkili meslek mensubunun sorumluluğunun yerine geçmez.")
    _madde(doc, "Yönetmelikte yer almayan mevzuatı (İSG, çevre, Seveso, ATEX vb.) denetlemez.")
    _not(doc, "Kılavuzda belirtildiği üzere, Kılavuz ile mevzuat arasında çelişme halinde yürürlükteki mevzuat geçerlidir. Yönetmelik değiştiğinde yangin/tablolar.py dosyasındaki tablolar güncellenmelidir.", "Önemli")

    # ---------------- 2 Akış
    doc.add_heading("2. Çalışma akışı: baştan sona 8 adım", 1)
    _tablo(doc, ["Adım", "Kim", "Ne yapılır", "Araç"], [
        ["1", "Mühendis", "Programdan 'Boş işveren formu'nu oluşturup işverene gönderir.", "Program / Excel formu"],
        ["2", "İşveren", "Formun sarı hücrelerini doldurur, belgelerin kopyalarını ekler (Belgeler sayfası).", "Excel formu"],
        ["3", "Mühendis", "Gerekirse sahada ölçüm yapar (çıkış genişliği, kaçış mesafesi, yükseklikler) ve teknik alanları (pompa, sprinkler hidroliği) ilgili firmadan alır.", "Ölçüm / belgeler"],
        ["4", "Mühendis", "'Formu içe aktar' ile doldurulmuş formu programa alır; uyarı ve eksik listesini inceler.", "Program"],
        ["5", "Mühendis", "Sekmeleri gözden geçirir, tehlike sınıfını ve varsayımları kontrol eder, eksikleri tamamlar.", "Program"],
        ["6", "Mühendis", "'Hesapla' düğmesine basar; 'Uygun değil', 'Gerekli', 'Veri girilmedi' satırlarını inceler.", "Program"],
        ["7", "Mühendis", "Düzeltme önerilerini işverene/müellife bildirir, gerekirse veriyi güncelleyip yeniden hesaplar.", "Program"],
        ["8", "Mühendis", "Excel ve Word raporunu kaydeder, projeyi .json olarak saklar.", "Program"],
    ], [1.2, 2.4, 10.4, 3.0])

    # ---------------- 3 Kurulum
    doc.add_heading("3. Kurulum", 1)
    doc.add_heading("3.1 Windows: YanginHesap.exe (önerilen, kurulum gerektirmez)", 2)
    _numara(doc, f"Tarayıcıda şu adresi açın: {RELEASE}")
    _numara(doc, "'YanginHesap.exe' dosyasını indirin ve istediğiniz klasöre koyun (ör. Masaüstü).")
    _numara(doc, "Çift tıklayın. İlk açılış birkaç saniye sürebilir (program kendini geçici klasöre açar).")
    _numara(doc, "Windows 'Bilinmeyen yayıncı' (SmartScreen) uyarısı verirse 'Ek bilgi' → 'Yine de çalıştır' seçin. Dosya dijital olarak imzalı olmadığı için bu uyarı normaldir.")
    _not(doc, "Bazı antivirüs programları imzasız .exe dosyalarını ilk çalıştırmada kontrol eder; gerekirse dosyayı güvenilir listeye ekleyin. Kaynak kodu açıktır: " + DEPO)
    doc.add_heading("3.2 Güncelleme", 2)
    _p(doc, f"Program her güncellendiğinde aynı sayfada ({RELEASE}) yeni sürüm yayınlanır. Eski .exe dosyasını silip yenisini indirmeniz yeterlidir; proje dosyalarınız (.json) etkilenmez.")
    doc.add_heading("3.3 Python ile çalıştırma (isteğe bağlı)", 2)
    _madde(doc, "Python 3.10+ kurun (python.org; kurulumda 'Add Python to PATH' kutusunu işaretleyin).")
    _madde(doc, "Depodaki 'calistir.bat' (Windows) veya 'calistir.sh' (Linux/macOS) dosyası web arayüzünü (Streamlit) kurar ve açar.")
    _madde(doc, "Masaüstü arayüzü için: pip install -r requirements-exe.txt ardından python masaustu.py.")
    doc.add_heading("3.4 Sistem gereksinimleri", 2)
    _tablo(doc, ["Öğe", "Gereksinim"], [["İşletim sistemi", "Windows 10/11 (64 bit) — .exe için"], ["Bellek / disk", "En az 4 GB RAM; ~100 MB disk"],
                                       ["Excel/Word", "Raporları açmak için Microsoft Office veya LibreOffice (isteğe bağlı; program dosyaları kendi üretir)"], ["İnternet", "Gerekmez (indirme dışında)"]], [4, 13])

    # ---------------- 4 Veri toplama
    doc.add_heading("4. İşverenden veri toplama (Excel formu)", 1)
    _p(doc, "Hesap için gereken bilgilerin büyük kısmı işletmede ve proje dosyalarında hazırdır. Bunları tek tek sormak yerine standart bir form kullanın: "
            "programda 'Boş işveren formu' düğmesi (web arayüzünde yan menüde 'Boş işveren formunu indir') Isveren_Veri_Toplama_Formu.xlsx dosyasını üretir.")
    doc.add_heading("4.1 Formun yapısı", 2)
    _tablo(doc, ["Sayfa", "İçerik", "Genelde kim doldurur"], [
        ["Talimat", "Formun nasıl doldurulacağı, renkler, ölçü birimleri", "—"],
        *[[s.ad, s.baslik.split(". ", 1)[-1] + f" ({len(s.alanlar)} soru)", ", ".join(sorted({A_KIM[a.kim] for a in s.alanlar}))] for s in A.SAYFALAR for A_KIM in [{"Z": "İşveren (zorunlu)", "V": "İşveren (varsa)", "M": "Mühendis/firma"}]],
        *[[t.ad, t.baslik.split("— ", 1)[-1], {"Z": "İşveren (zorunlu)", "V": "İşveren (varsa)", "M": "Mühendis/firma"}[t.kim]] for t in A.TABLO_SAYFALARI],
        ["Belgeler", "İstenecek belgelerin kontrol listesi (17 belge)", "İşveren"],
        ["Notlar", "Forma sığmayan ek bilgiler", "İşveren"],
    ], [3.4, 9.6, 4.0])
    doc.add_heading("4.2 Hücre renkleri ve 'Kim doldurur' etiketi", 2)
    _tablo(doc, ["Görünüm", "Anlamı"], [
        ["SARI hücre", "Cevap yazılacak yer. Yalnızca sarı hücrelere yazın."],
        ["Turuncu 'İŞVEREN — zorunlu'", "Bilgi işverenden mutlaka beklenir. Boş bırakılırsa program 'eksik zorunlu bilgi' olarak bildirir."],
        ["Sarımsı 'İŞVEREN — varsa'", "İşletmede varsa/biliniyorsa doldurulur; yoksa boş bırakılır."],
        ["Mavi 'MÜHENDİS / montaj firması'", "Teknik değer; proje müellifi, montaj/bakım firması veya mühendis doldurur. İşveren ilgili firmaya yönlendirebilir."],
        ["Gri satır (tablolarda 6. satır)", "ÖRNEK satırdır; silinmez, programa aktarılmaz."],
        ["Açılır liste", "Listeden seçin; listede olmayan değer yazılamaz."],
    ], [5, 12])
    doc.add_heading("4.3 İşverene gönderirken", 2)
    _madde(doc, "Formla birlikte kısa bir e-posta ekleyin: son teslim tarihi, kime ulaşılacağı, belge kopyalarının (PDF/fotoğraf) formla birlikte gönderilmesi.")
    _madde(doc, "'Bilmiyorsanız boş bırakın, tahmin yürütmeyin' deyin. Yanlış veri, boş veriden daha zararlıdır; boş bırakılan alan raporda 'VERİ GİRİLMEDİ' olarak görünür ve kapatılabilir.")
    _madde(doc, "Sayfa/sütun silinmemeli, sayfa adları değiştirilmemelidir; aksi halde içe aktarma o sayfayı okuyamaz (program uyarı verir).")
    doc.add_heading("4.4 İşverenden istenecek belgeler", 2)
    _tablo(doc, ["No", "Belge", "Neden gerekli", "Kimden", "Zorunluluk"], [[str(i + 1), b, nd, k, z] for i, (b, nd, k, z) in enumerate(A.BELGELER)], [0.9, 5.6, 5.7, 3.0, 1.8], boyut=8)
    doc.add_heading("4.5 Sahada doğru ölçüm için ipuçları", 2)
    for baslik, metin in [
        ("Temiz genişlik: ", "Kapı için kanat 90° açıkken kullanılabilir boşluk; merdiven için küpeştenin çıkıntısının 80 mm'si dahil olmak üzere kullanılabilir genişlik. Kasa/çerçeve ölçüsü değildir."),
        ("Kaçış mesafesi: ", "Mekânın en uzak noktasından (duvarlardan 40 cm içeride) en yakın çıkışa, makine ve raf engellerini dolanarak, gerçek yürüme yolu boyunca ölçülür; kuş uçuşu değildir (kuş uçuşu ayrı bir alandır)."),
        ("Yükseklikler: ", "Yapı yüksekliği bodrum ve çatı arası dahil tüm katların toplamıdır; bina yüksekliği zemin yaklaşma kotundan saçak seviyesine kadardır. İkisini karıştırmayın."),
        ("Alanlar: ", "Kullanıcı yükü için yemekhane/toplantı/sergi türlerinde NET alan, diğerlerinde BRÜT alan kullanılır (Ek-5/A)."),
        ("Çalışan sayısı: ", "En kalabalık vardiya + sürekli bulunan taşeron/ziyaretçi. Vardiya değişiminde üst üste binen sayıyı dikkate alın."),
        ("Yanıcı madde: ", "Üretimde kullanılan ve depolanan tüm yanıcı/parlayıcı maddeleri (boya, tiner, solvent, yapıştırıcı, köpük, plastik, ahşap tozu, LPG, akaryakıt, ambalaj) MSDS'ten kontrol edin; hiçbirini atlamayın."),
        ("Pompa verileri: ", "Pompa etiketi ve karakteristik eğrisinden okunur; eğri üzerindeki üç nokta: nominal debi, sıfır debi (kapalı vana), %150 debi."),
    ]:
        _madde(doc, metin, baslik)

    # ---------------- 5 Ekran
    doc.add_heading("5. Program ekranı: düğmeler ve sekmeler", 1)
    _resim(doc, ekran("01_genel.png"), "Şekil 1 — Ana pencere: üstte düğme çubuğu, altında sekmeler (Örnek fabrika yüklenmiş hali).")
    doc.add_heading("5.1 Düğme çubuğu", 2)
    _tablo(doc, ["Düğme", "İşlevi"], [
        ["▶ Hesapla", "Girilen tüm verilerle hesaplamayı çalıştırır ve 'Sonuçlar' sekmesine geçer."],
        ["Excel raporu", "Hesaplar ve sonucu .xlsx olarak kaydeder (özet, tüm sonuçlar, eksikler, kullanıcı yükü, dayanak sayfaları)."],
        ["Word raporu", "Hesaplar ve sonucu .docx olarak kaydeder (proje bilgileri, özet, uygun olmayan hususlar, bölüm bölüm tablolar, dayanak)."],
        ["Projeyi aç / Projeyi kaydet", "Girdileri .json dosyası olarak saklar/geri yükler. Sonradan düzeltme ve yeniden hesaplama için kullanın."],
        ["Boş işveren formu", "İşverene verilecek Excel veri toplama formunu kaydeder."],
        ["Formu içe aktar", "Doldurulmuş formu okur, alanları ve tabloları doldurur; eksik ve hatalı alanları listeler."],
        ["Örnek fabrika", "Çözümlü örnek projeyi yükler (programı tanımak için)."],
        ["Sıfırla", "Tüm alanları varsayılan değerlere döndürür. Kaydetmediğiniz veriler kaybolur."],
    ], [4.5, 12.5])
    doc.add_heading("5.2 Sekmeler", 2)
    _tablo(doc, ["Sekme", "Ne girilir"], [
        ["Genel", "Proje bilgileri, sınıflandırma, aktif sistemler, bina geometrisi, betonarme/ahşap detayları"],
        ["Kullanıcı yükü ve kaçış", "Mahaller tablosu (kullanıcı yükü) ve Kat kaçış verileri tablosu (çıkış genişliği, mesafeler)"],
        ["Merdiven, kapı, duman, tank", "Kaçış merdivenleri, kapılar, mekanik duman tahliyesi mahalleri, yanıcı sıvı tankları"],
        ["Su ve söndürme", "Sprinkler, dolap/hidrant ve yangın pompası bilgileri"],
        ["Algılama, duman, aydınlatma", "Buton, acil aydınlatma/yönlendirme ölçümleri, basınçlandırma"],
        ["Kazan ve tehlikeli madde", "Kazan dairesi, yakıt, trafo/jeneratör, LPG, yanıcı sıvı depolama"],
        ["Çevre ve ekipler", "İtfaiye erişimi, orman alanı, çalışan sayısı"],
        ["Sonuçlar", "Özet, durum filtresi ve tüm sonuçların renkli tablosu"],
    ], [5, 12])
    doc.add_heading("5.3 Tablolarla çalışmak", 2)
    _madde(doc, "'Satır ekle' yeni satır açar; 'Düzenle' (veya satıra çift tık) seçili satırı düzenleme penceresinde açar; 'Sil' seçili satırı kaldırır; 'Kopyala' satırı çoğaltır.")
    _madde(doc, "Sütun başlıkları uzunsa yatay kaydırma çubuğunu kullanın.")
    _madde(doc, "Kat adları iki tabloda (Mahaller ve Kat kaçış verileri) aynı yazılmalıdır; aksi halde kat için çıkış verisi eşleşmez.")
    _madde(doc, "Sayıları virgül veya nokta ile yazabilirsiniz (12,5 veya 12.5).")

    # ---------------- 6 Veri girişi
    doc.add_heading("6. Verilerin girilmesi: sekme sekme, alan alan", 1)
    _p(doc, "Bu bölümdeki tablolar, işveren formundaki sorularla birebir aynıdır. 'Kim' sütunu: İ = işveren (zorunlu), İv = işveren (varsa), M = mühendis/firma. "
            "'Programda etkisi' sütunu alanın hangi hesaba girdiğini ve yönetmelik dayanağını gösterir.")
    ek_resim = {"1_Genel": "01_genel.png", "3_Mevcut_Sistemler": "02_su.png", "4_Aydinlatma_Duman": "03_algilama.png", "5_Kazan_Yakit_LPG": "04_kazan.png", "6_Cevre_Erisim": "05_cevre.png"}
    kisalt = {"Z": "İ", "V": "İv", "M": "M"}
    for no, s in enumerate(A.SAYFALAR, 1):
        doc.add_heading(f"6.{no} {s.baslik.split('. ', 1)[-1]}  (Program sekmesi: {s.program_sekmesi})", 2)
        _p(doc, s.giris)
        if s.ad in ek_resim:
            _resim(doc, ekran(ek_resim[s.ad]), f"Şekil — {s.program_sekmesi} sekmesi", 15.5)
        _tablo(doc, ["Alan", "Açıklama / nasıl doldurulur", "Birim", "Kim", "Kaynak", "Programda etkisi"],
               [[a.baslik, a.aciklama, a.birim, kisalt[a.kim], a.kaynak, a.madde] for a in s.alanlar],
               [3.9, 4.7, 1.1, 0.8, 2.6, 3.9], boyut=7.5)
    doc.add_heading("6.8 Tablolar", 2)
    for t in A.TABLO_SAYFALARI:
        doc.add_heading(t.baslik, 3)
        _p(doc, t.giris)
        _p(doc, f"Program sekmesi: {t.program_sekmesi}.  Kaynak: {t.kaynak}.", italik=True, boyut=9)
        kolonlar = E.TABLOLAR[t.anahtar]
        _tablo(doc, ["Sütun", "Açıklama"], [[b, t.aciklamalar.get(k, "")] for k, b, _t, _s, _g in kolonlar], [4.6, 12.4], boyut=8)
        if t.anahtar == "mahal":
            _resim(doc, ekran("06_kacis.png"), "Şekil — 'Kullanıcı yükü ve kaçış' sekmesi: Mahaller ve Kat kaçış verileri tabloları", 15.5)
        if t.anahtar == "merdiven":
            _resim(doc, ekran("07_merdiven.png"), "Şekil — 'Merdiven, kapı, duman, tank' sekmesi", 15.5)
    doc.add_heading("6.9 Ek-5/A kullanıcı yükü katsayıları (Mahal türleri)", 2)
    _tablo(doc, ["Mahal türü", "m²/kişi", "Alan türü"], [[a, f"{k:g}", t] for a, k, t in T.EK5A if k > 0], [11, 2.5, 3.5], boyut=8)
    doc.add_heading("6.10 Tehlike sınıfı için faaliyet listesi (Ek-1)", 2)
    _p(doc, "'Üretim/tesis türü' listesinden seçilen faaliyetin tehlike sınıfı otomatik önerilir. Birden fazla faaliyet varsa en yüksek sınıf esas alınır (Madde 19(1)). "
            "Orta Tehlike-1/2 faaliyetlerde boyama işlemi gibi yüksek yangın yükü olan alanlar varsa Orta Tehlike-3 seçilmelidir (Ek-1/B dipnotu).")
    _tablo(doc, ["Tehlike sınıfı", "Faaliyetler"], [[T.TEHLIKE_SINIFLARI[kod], "; ".join(f"{ad}" for _k, ad, sn in T.EK1_TESISLER if sn == kod)] for kod in T.TEHLIKE_SINIFLARI if any(sn == kod for _k, _a, sn in T.EK1_TESISLER)],
           [3.3, 13.7], boyut=7.5)

    # ---------------- 7 Sonuçlar
    doc.add_heading("7. Hesaplama ve sonuçların okunması", 1)
    _p(doc, "'Hesapla' düğmesine bastığınızda program tüm modülleri çalıştırır; 'Sonuçlar' sekmesi açılır. Üstte durum sayıları, altında renkli sonuç tablosu bulunur.")
    _resim(doc, ekran("08_sonuclar.png"), "Şekil — Sonuçlar sekmesi: durum sayıları ve renkli sonuç tablosu")
    doc.add_heading("7.1 Durum etiketleri", 2)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for j, b in enumerate(["Durum", "Anlamı", "Ne yapmalı"]):
        _hucre(t.rows[0].cells[j], b, kalin=True, renk=RGBColor(0xFF, 0xFF, 0xFF))
        _golge(t.rows[0].cells[j], "1F3864")
    for ad, renk, anlam, eylem in DURUM_ACIKLAMA:
        c = t.add_row().cells
        _hucre(c[0], ad, kalin=True)
        _golge(c[0], renk)
        _hucre(c[1], anlam)
        _hucre(c[2], eylem)
    for row in t.rows:
        for j, w in enumerate([3.2, 7.0, 6.8]):
            row.cells[j].width = Cm(w)
    doc.add_paragraph()
    doc.add_heading("7.2 Tablo sütunları", 2)
    _tablo(doc, ["Sütun", "Anlamı"], [["Bölüm", "Hesabın ait olduğu modül (alt başlıkta kat/merdiven adı)"], ["Kalem", "Kontrol edilen veya hesaplanan büyüklük"],
                                      ["Değer", "Sonuç veya girilen değer (birimiyle)"], ["Durum", "Yukarıdaki durum etiketleri"], ["Madde / Ek", "Dayanak yönetmelik maddesi veya eki"],
                                      ["Açıklama", "Formül, gerekli değer, koşul ve varsayımlar. Satıra çift tıklayarak tam metni okuyun."]], [3.5, 13.5])
    doc.add_heading("7.3 Filtre", 2)
    _p(doc, "'Durum filtresi' listesinden yalnızca 'Uygun değil', 'Gerekli', 'Veri girilmedi' vb. satırları gösterebilirsiniz. Varsayılan görünüm 'Bilgi hariç hepsi'dir; ara hesapları da görmek için 'Hepsi' seçin.")
    _resim(doc, ekran("09_uygun_degil.png"), "Şekil — Filtre: yalnızca 'UYGUN DEĞİL' satırları")
    doc.add_heading("7.4 'Uygun değil' sonuçlarını giderme yolları", 2)
    _tablo(doc, ["Tipik uygunsuzluk", "Olası düzeltmeler"], [
        ["Çıkış genişliği yetersiz", "Çıkış kapısı/merdiven genişliğini artırın; yeni çıkış ekleyin; mekândaki kişi sayısını azaltın (belirli kişi sayısı/yerleşim); bölümleri ayırın."],
        ["Tekil çıkış asgari genişliğinin altında (≥100/150/200 cm)", "Dar kapıyı genişletin veya çıkış sayısını artırarak yükü bölün; asgari tekil genişlik kullanıcı sayısına bağlıdır (Madde 33)."],
        ["Kaçış mesafesi aşılıyor", "Ek çıkış açın; sprinkler ekleyin (limit artar); iç bölme/raf yerleşimini değiştirerek yolu kısaltın; kolay alevlenici madde yoksa endüstriyel çarpanın uygulanıp uygulanmadığını kontrol edin."],
        ["Kompartıman alanı aşılıyor", "Yangın duvarı ile bölün; otomatik algılama/sprinkler/duman tahliye ekleyin (Ek-4 dipnotları)."],
        ["Yağmurlama zorunlu ama yok", "Madde 96(2) kapsamı (ör. kolay alevlenici madde ve 1000 m² üstü) → sprinkler projesi."],
        ["Yangın dolabı/hidrant zorunlu ama yok", "Dolap sistemi / dış hidrant projelendirin (Madde 94-95)."],
        ["Su deposu yetersiz", "Depoyu büyütün; tasarım yoğunluğu/koruma alanını azaltacak hidrolik optimizasyon (mühendis); Madde 92(4) hidrolik hesap."],
        ["Pompa karakteristiği uygunsuz", "Kapalı vana basma ≤ %140, %150 debide ≥ %65 koşullarına uygun pompa seçin (Madde 93)."],
        ["Merdiven ölçüleri (rıht/basış/sahanlık)", "Merdiven geometrisini düzeltin (rıht ≤ 175 mm, basış ≥ 250 mm, sahanlık ≤ 17 basamak)."],
        ["Merdiven/kapı yangın dayanımı yetersiz", "Daha yüksek dayanımlı duman sızdırmaz kapı (60/90 dk) ve 120 dk duvar."],
    ], [5.5, 11.5])

    # ---------------- 8 Raporlar
    doc.add_heading("8. Raporlar (Excel ve Word) ve proje dosyası", 1)
    doc.add_heading("8.1 Excel raporu (.xlsx)", 2)
    _tablo(doc, ["Sayfa", "İçerik"], [["Özet", "Proje bilgileri, durum sayıları, uyarı"], ["Sonuçlar", "Tüm sonuç satırları (filtrelenebilir, renkli)"], ["Eksikler", "'Uygun değil' ve 'Veri girilmedi' satırları (işverene gönderilecek liste)"],
                                       ["Kullanıcı yükü", "Kat ve mahal bazında kişi hesabı"], ["Dayanak", "Kaynak mevzuat, Danıştay kararı notu, uyarı"]], [3.5, 13.5])
    doc.add_heading("8.2 Word raporu (.docx)", 2)
    _p(doc, "Proje bilgileri → sonuç özeti → uygun olmayan hususlar ve girilmeyen veriler → bölüm bölüm hesap tabloları (kalem, değer, durum, madde/açıklama) → dayanak ve uyarılar. "
            "Rapor dosyasını kaydettiğinizde program açmak isteyip istemediğinizi sorar.")
    doc.add_heading("8.3 Proje dosyası (.json)", 2)
    _p(doc, "'Projeyi kaydet' tüm girdileri .json olarak saklar. Yeniden açmak, veriyi güncelleyip tekrar hesaplamak ve farklı senaryoları (ör. sprinkler var/yok) karşılaştırmak için kullanın. "
            "Dosya düz metindir; sürüm yönetimine ve yedeklemeye uygundur.")
    _not(doc, "Rapor tarihi 'Tarih' alanından alınır; boş bırakılırsa raporlama günü yazılır. 'Hazırlayan' alanına raporu hazırlayan mühendisin adını yazın.")

    # ---------------- 9 Hesap yöntemleri
    doc.add_heading("9. Hesap yöntemleri ve formüller", 1)
    _p(doc, "Aşağıda programın kullandığı başlıca formüller ve kurallar özetlenmiştir. Her satırın tam gerekçesi raporun 'Açıklama' sütununda yer alır.")
    yontemler = [
        ("Kullanıcı yükü (Madde 32, Ek-5/A)", "N = Σ (mahal alanı / katsayı). Kişi sayısı belirliyse N = max(hesaplanan, belirli sayı). Aynı anda kullanılmayan mahaller (tuvalet, soyunma, depo) hesaba katılmayabilir (Madde 31(6))."),
        ("Gerekli çıkış genişliği (Madde 32(2), 33(1), Ek-5/B)", "W = (N / birim genişlik için kişi sayısı) × 50 cm. Birim genişlik değerleri Ek-5/B'dedir (endüstriyel yapı: dış kapı 100, diğer kapı 80, merdiven 60, koridor 100 kişi)."),
        ("Asgari tekil genişlik (Madde 33)", "50–500 kişi: ≥100 cm; 501–2000: ≥150 cm; ≥2001: ≥200 cm; her durumda ≥80 cm; koridor/hol ≥110 cm; yüksek binada ≥120 cm. İki çıkışlı mekânda her çıkış ≥ W/2."),
        ("Çıkış sayısı (Madde 39, 52)", "Aksi belirtilmedikçe en az 2; >500 kişi en az 3; >1000 kişi en az 4. Fabrika/depo/büroda tek çıkış yalnızca yapı yüksekliği <21,50 m, kat kullanıcısı <50, mesafe Ek-5/B'ye uygun, yanmaz yapı ve kolay alevlenici madde yok ise."),
        ("Kaçış uzaklığı (Madde 32(3)-(6), Ek-5/B)", "Tablo sınırı kullanım sınıfı, yön (tek/iki) ve sprinkler varlığına göre. Endüstriyel yapıda kolay alevlenici üretim yoksa ×1,5 (mevcut yapıda ×2, Ek-14). Kuş uçuşu mesafe ≤ sınırın 2/3'ü; çıkmaz koridor ≤15/20 m; iki çıkış arası ≥ diyagonal/2 (sprinklerli: /3)."),
        ("Yangın dayanım süresi (Ek-3/C)", "Satır: kullanım sınıfı + sprinkler; sütun: bina yüksekliği (≤21,50 / ≤30,50 / >30,50 m) ve bodrum derinliği (>10 / 5-10 / <5 m). 'İzin verilmez' hücrelerinde yapı yapılamaz."),
        ("Kompartıman alanı (Ek-4)", "Endüstriyel OT-1/2: 15 000 m² (tek katlı sınırsız); OT-3+: 6 000 m²; depo OT-1/2: 5 000 m², OT-3+: 1 000 m². Uygun yangın kontrol sistemleri (algılama, sprinkler, duman tahliye) varsa sınırsız (dipnot 2/3)."),
        ("Sprinkler debisi (Ek-8/B)", "Q = tasarım yoğunluğu (mm/dk) × koruma alanı (m²) [l/dk]. Örnek: OT-3 ıslak 5 mm/dk × 216 m² = 1 080 l/dk."),
        ("Su deposu (Madde 92)", "V = (Q_sprinkler + Q_dolap + Q_hidrant) × süre / 1000 [m³]; süre: düşük 30, orta 60, yüksek 90 dk. Q_dolap = eş zamanlı dolap sayısı × Ek-8/C debisi; Q_hidrant: Ek-8/C. Ek-8/A tablo değeri ile karşılaştırılır; büyük olan önerilir. Yalnız hidrant: ≥1900 l/dk × 90 dk = 171 m³."),
        ("Yangın pompası (Madde 93)", "Q_pompa (m³/h) = Q_toplam (l/dk) × 0,06. Hy = statik yükseklik + boru kaybı + akma basıncı (hidrant 700 kPa ≈ 70 mSS). Kapalı vana basma ≤ 1,4 × anma; %150 debide ≥ 0,65 × anma; pompa anma debisinin %130'una kadar yük karşılar; yedek pompa."),
        ("Yangın dolabı, hidrant (Madde 94-95)", "Dolap aralığı ≤30 m (sprinklerli ve su alma ağzı varsa 45 m); adet = ⌈boy/aralık⌉ × ⌈en/aralık⌉ × kat. Hidrant aralığı 50/100/125/150 m (bölge riskine göre); adet = ⌈bina çevresi/aralık⌉; çıkış basıncı 700 kPa; debi ≥1900 l/dk."),
        ("Söndürücü (Madde 99)", "Düşük tehlike 500 m², orta/yüksek 250 m² yapı alanı başına 1 adet 6 kg; her noktadan en yakın cihaza ≤25 m (kare ızgara kapsaması); büyük olan alınır."),
        ("Yangın uyarı butonu (Madde 75)", "Her noktadan en yakın butona yatay erişim ≤60 m (kare ızgara: kenar = 60·√2 ≈ 85 m); her çıkış yanında bir buton; zorunluluk: 2–4 katlı ve kat alanı >400 m², >4 kat veya yüksek bina."),
        ("Acil aydınlatma (Madde 72-73)", "Süre 60 dk (kullanıcı >200 ise 120 dk); merkez hatta ≥1 lux, süre sonunda ≥0,5 lux, max/min ≤40. Yönlendirme levhası görülebilirlik uzaklığı = yükseklik × 200 (içten) veya × 100 (dıştan); işaret adedi = ⌈yol uzunluğu / uzaklık⌉."),
        ("Mekanik duman tahliyesi (Madde 88(3))", "Kazan dairesi, kapalı otopark, bodrum depo alanı >2000 m² ise zorunlu; debi = 10 × hacim [m³/h]."),
        ("Basınçlandırma (Madde 89)", "Konut dışı merdiven kovası >30,50 m veya bodrum >4 ise zorunlu. Asgari hava debisi Q = 1 m/s × kapı alanı × 3 açık kapı + sızıntı; basınç farkı ≥50 Pa (kapı açıkken ≥15 Pa)."),
        ("Orman alanı (Madde 21(5))", "Dış yangın bölgesi 100 m; eğim %30–55: aşağı yönde ×2, diğer yönlerde ×1,5; >%55: ×4 ve ×2."),
        ("Acil durum ekipleri (Madde 126)", ">50 kişi: söndürme ≥3, kurtarma ≥3, koruma ≥2, ilk yardım ≥2 (toplam ≥10 kişi); yılda en az 1 tatbikat."),
        ("Yanıcı sıvı (Madde 114, 118-119)", "Sınıf IA cinsinden toplam = IA + IB/2 + IC/4 + II/12 + IIIA/40 + IIIB/80 ≤ 12 500 L; Ek-11 eşikleri (bildirim/izin); havuzlama ≥ en büyük tank; tank mesafeleri Ek-12/C-Ç."),
    ]
    _tablo(doc, ["Konu", "Yöntem / formül"], [[a, b] for a, b in yontemler], [4.6, 12.4], boyut=8)

    # ---------------- 10 Çözümlü örnek
    g = ornek_fabrika()
    s = H.hesapla(g)
    doc.add_heading("10. Çözümlü örnek: Örnek Mobilya Fabrikası", 1)
    _p(doc, "Programın 'Örnek fabrika' düğmesi bu projeyi yükler. Örnek, kaçış verilerinde bilerek iki uygunsuzluk içerir; amaç sonuçların nasıl okunup düzeltileceğini göstermektir.")
    doc.add_heading("10.1 Girdiler (özet)", 2)
    _tablo(doc, ["Girdi", "Değer"], [
        ["Faaliyet / sınıf", "Ahşap/mobilya fabrikası → Orta Tehlike-3; endüstriyel yapı; kolay alevlenici madde var (tiner)"],
        ["Yükseklik, kat", f"Yapı {g.yapi_yuksekligi:g} m, bina {g.bina_yuksekligi:g} m; {g.kat_sayisi} kat; bodrum yok"],
        ["Alanlar", f"Taban {g.taban_alani:,.0f} m²; toplam kapalı {g.toplam_kapali_alan:,.0f} m²; en büyük kat {g.en_buyuk_kat_alani:,.0f} m²".replace(",", ".")],
        ["Sistemler", "Islak sprinkler, algılama, duman tahliye, hidrant, dolap; çelik taşıyıcı"],
        ["Zemin kat mahalleri", "Üretim holü 4800 m² (10 m²/kişi), hammadde deposu 900 m² (30 m²/kişi), yemekhane 150 m² net (1,5 m²/kişi)"],
        ["Zemin kat çıkışları", "4 dış kapı, toplam 480 cm, en dar 120 cm; en uzak mesafe 70 m (iki yön)"],
        ["1. kat", "Ofis 900 m² (10 m²/kişi); 2 kaçış merdiveni toplam 300 cm"],
        ["Pompa / depo", "Pompa 150 m³/h, 110 mSS (kapalı vana 140, %150 debide 75 mSS); su deposu 200 m³; statik 14 + kayıp 18 + akma 70 mSS"],
    ], [4.2, 12.8])
    doc.add_heading("10.2 Adım adım sonuçlar", 2)
    yuk = _bul(s, "Kat kullanıcı yükü", "Zemin")
    gen = _bul(s, "Gerekli toplam genişlik — Dışarı", "Zemin")
    dar = _bul(s, "En dar tekil", "Zemin")
    msf = _bul(s, "Ölçülen en uzak", "Zemin")
    lim = _bul(s, "En çok kaçış uzaklığı", "Zemin")
    qs = _bul(s, "Yağmurlama debisi")
    qt = _bul(s, "Toplam yangın suyu debisi")
    vd = _bul(s, "Hesaplanan su deposu")
    hy = _bul(s, "basma yüksekliği Hy")
    ek8a = _bul(s, "Ek-8/A")
    day = _bul(s, "Taşıyıcı sistem / kompartıman")
    komp = _bul(s, "Tablo sınırı (Ek-4)")
    dolap = _bul(s, "Yaklaşık yangın dolabı")
    sond = _bul(s, "Gerekli en az söndürücü")
    _tablo(doc, ["Adım", "Hesap", "Sonuç"], [
        ["1 Kullanıcı yükü (zemin)", "4800/10 + 900/30 + 150/1,5 = 480 + 30 + 100", f"{_fmt(yuk.deger)} kişi"],
        ["2 Gerekli çıkış genişliği (dış kapı, birim 100 kişi)", f"{_fmt(yuk.deger)} / 100 × 50 cm", f"{_fmt(gen.deger)} cm  → mevcut 480 cm: UYGUN"],
        ["3 Tekil çıkış asgari genişliği", ">500 kişi → her çıkış ≥150 cm", f"En dar çıkış {_fmt(dar.deger)} cm → {dar.durum}"],
        ["4 Kaçış mesafesi sınırı", "Endüstriyel, iki yön, sprinklerli: 60 m; kolay alevlenici madde var → çarpan uygulanmaz", f"Sınır {_fmt(lim.deger)} m; ölçülen {_fmt(msf.deger)} m → {msf.durum}"],
        ["5 Yangın dayanımı (Ek-3/C)", "Endüstriyel, sprinklerli, bina yüksekliği 10 m (<21,50)", f"{_fmt(day.deger)} dk"],
        ["6 Kompartıman (Ek-4)", "OT-3 → 6000 m²; sprinkler/algılama/duman tahliye var → sınırsız (koşullu)", f"Tablo sınırı {_fmt(komp.deger)} m²"],
        ["7 Sprinkler debisi", "OT-3 ıslak: 5 mm/dk × 216 m²", f"{_fmt(qs.deger)} l/dk"],
        ["8 Toplam debi", "1080 + 2 dolap × 100 + hidrant 1000 (OT-3/4)", f"{_fmt(qt.deger)} l/dk"],
        ["9 Su deposu", f"{_fmt(qt.deger)} l/dk × 60 dk / 1000", f"{_fmt(vd.deger)} m³ (Ek-8/A: {_fmt(ek8a.deger)} m³) → mevcut 200 m³: UYGUN"],
        ["10 Pompa Hy", "14 + 18 + 70", f"{_fmt(hy.deger)} mSS → pompa 110 mSS ≥ Hy: UYGUN"],
        ["11 Dolap ve söndürücü", "Izgara hesabı / 250 m² başına 1", f"≈ {_fmt(dolap.deger)} dolap; ≥ {_fmt(sond.deger)} söndürücü"],
    ], [4.2, 7.0, 5.8], boyut=8)
    _not(doc, "Pompa kontrolünde üç koşul birlikte aranır: anma basma yüksekliği ≥ Hy; kapalı vana basma ≤ %140 × anma (140 ≤ 154 mSS); %150 debide basma ≥ %65 × anma (75 ≥ 71,5 mSS).", "Pompa")
    doc.add_heading("10.3 Uygun olmayan iki husus ve düzeltme", 2)
    _tablo(doc, ["Sonuç", "Neden", "Düzeltme"], [
        [f"En dar tekil çıkış {_fmt(dar.deger)} cm", "Zemin katta 610 kişi var (>500); tek bir çıkış en az 150 cm olmalı (Madde 33(1))", "120 cm'lik çıkışı en az 150 cm'e genişletin veya çıkış sayısını artırıp tekil genişliği yeniden dağıtın."],
        [f"Kaçış mesafesi {_fmt(msf.deger)} m > {_fmt(lim.deger)} m", "Kolay alevlenici madde bulunduğu için Ek-5/B'deki +%50 artış uygulanamaz", "Zemin kata ek çıkış açın, yerleşimi değiştirin veya kolay alevlenici madde yoksa yok olarak gösterin (doğruysa); sonra yeniden hesaplayın."],
    ], [4.2, 6.4, 6.4], boyut=8)
    _p(doc, "Düzeltilen değerler (ör. tekil çıkış 150 cm, en uzak mesafe 58 m) forma/programa girilip 'Hesapla' tekrar basıldığında bu iki satır 'UYGUN'a döner. "
            "Örnek, 'Örnek fabrika' düğmesiyle yüklenip değerler değiştirilerek denenebilir.")

    # ---------------- 11 SSS
    doc.add_heading("11. Sorun giderme ve sık sorulan sorular", 1)
    sss = [
        ("Program açılmıyor / hiçbir şey olmuyor.", "İlk açılış birkaç saniye sürer. Antivirüs/SmartScreen engeli olabilir (3.1). Dosyayı farklı bir klasöre (ör. Masaüstü) kopyalayıp yeniden deneyin."),
        ("'Formu içe aktar' uyarı veriyor: '... sayfası bulunamadı'.", "Formdaki sayfa adı değiştirilmiş veya sayfa silinmiştir. Orijinal şablondan yeniden doldurun; sayfa adlarını (1_Genel, T1_Mahaller ...) değiştirmeyin."),
        ("'... alanındaki değer anlaşılamadı' uyarısı.", "Sayı alanına metin yazılmış (ör. 'yüz metre'). İşverenden düzeltmesini isteyin veya programda ilgili alanı elle girin. Sayı alanlarına birim yazılmamalıdır."),
        ("'T1_Mahaller'deki ... katı T2_Katlar tablosunda yok' uyarısı.", "Kat adları iki tabloda aynı yazılmamış (Zemin kat / Zemin Kat farkı da önemlidir). Aynı yazın."),
        ("Çok sayıda 'VERİ GİRİLMEDİ' var.", "Normaldir; ölçülü değerler boşsa program gerekli değeri hesaplar ama kontrol edemez. Eksik değerleri işverenden isteyin (Eksikler sayfası bu iş için hazırlanmıştır)."),
        ("Tehlike sınıfını nasıl seçeceğim?", "Genel sekmesinde 'Üretim/tesis türü' listesinden faaliyeti seçin; sınıf otomatik gelir. Karışık faaliyetlerde en yüksek sınıfı seçin. Listede olmayan faaliyet için yönetmeliğin benzer faaliyetine göre mühendis kararı gerekir."),
        ("Ek-7 neden algılama zorunluluğu göstermiyor?", "Ek-7'nin endüstriyel, ticaret, kurum ve toplanma yapıları satırları Danıştay kararıyla iptal edildi (E.2019/261, K.2021/5537; İDDK onama 22/2/2023). Program bu hali esas alır; fabrikalarda algılama gerekliliği Madde 81, Ek-4 dipnotları ve itfaiye görüşüyle değerlendirilir."),
        ("Sprinkler yokken sprinkler zorunlu çıkıyor.", "Madde 96(2): yapı yüksekliği >30,50 m (konut dışı), kapalı otopark >600 m², kolay alevlenici/parlayıcı madde üretilen/bulundurulan yapı >1000 m², ticaret >2000 m² vb. Zorunluluk koşulunun kapsamında olup olmadığınızı Genel sekmesindeki 'kolay alevlenici' kutusunu kontrol ederek doğrulayın."),
        ("Sayıları virgülle mi noktayla mı yazmalıyım?", "İkisi de kabul edilir (12,5 / 12.5). Binlik ayırıcı kullanmayın."),
        ("Raporda Türkçe karakterler bozuk.", "Excel/Word dosyaları Unicode üretilir; farklı bir programla açıyorsanız Microsoft Office veya LibreOffice kullanın."),
        ("Sonuçlar projemle farklı, hangisi doğru?", "Program ön hesap ve yönetmelik kontrolüdür. Detay tasarım (hidrolik, yapısal) ve itfaiye görüşü ile birlikte değerlendirin; farkı rapordaki 'Açıklama' sütunundan inceleyin."),
        ("Yönetmelik değişirse?", "Tablolar yangin/tablolar.py içindedir; güncel tabloları işleyip uygulamayı yeniden derleyin (depodaki GitHub Actions iş akışı exe'yi otomatik üretir)."),
        ("Verilerim kaydoldu mu?", "Program verileri otomatik saklamaz. Çalışmayı 'Projeyi kaydet' ile .json olarak saklayın."),
    ]
    for q, a in sss:
        p = doc.add_paragraph()
        p.add_run("S: " + q).bold = True
        p.paragraph_format.space_after = Pt(0)
        p2 = doc.add_paragraph("C: " + a)
        p2.paragraph_format.space_after = Pt(6)

    # ---------------- 12 Sınırlar
    doc.add_heading("12. Sınırlar, varsayımlar ve sorumluluk", 1)
    for b, m in [
        ("Ön hesap değerleri: ", "Yaklaşık başlık, dolap, hidrant, buton ve söndürücü adetleri, yönetmeliğin azami uzaklık/alan kriterlerinden ızgara kapsamasıyla türetilir. Kesin yerleşim proje çizimiyle yapılır."),
        ("Yönetmelikte olmayıp varsayım gerektirenler: ", "Hidrant bölge riski (Madde 95(3) tehlike sınıfı eşleştirmesi vermez), OT-2 ve üzeri için başlık başına koruma alanı (TS EN 12845: 12 m² / 9 m²), eş zamanlı dolap sayısı (Kılavuz örneği: 2). Bunlar arayüzde değiştirilebilir."),
        ("Standart tabanlı detaylar kapsam dışıdır: ", "Boru çapı/hidrolik, duman tahliye hacim hesabı, sızıntı alanından basınçlandırma debisi, yapısal yangın analizi."),
        ("Mevcut yapılar: ", "Onuncu Kısım (Madde 138-167) hükümlerinden yalnızca Ek-14 kaçış uzaklığı çarpanı uygulanır; diğer uyum hükümleri ayrıca değerlendirilmelidir."),
        ("Tablo okuma güvencesi: ", "Tüm tablolar PDF'teki tablo görüntüleriyle karşılaştırılarak girilmiştir; testler Kılavuzun çözümlü örneklerini üretir (su deposu 79,2 m³; çıkış genişlikleri 168,87/101,32 cm)."),
        ("Sorumluluk: ", "Rapor, projenin onayı veya itfaiye görüşünün yerine geçmez. Sonuçların kullanımı ve projeye yansıtılması yetkili meslek mensubunun sorumluluğundadır. Mevzuatla çelişme halinde yürürlükteki mevzuat geçerlidir."),
    ]:
        _madde(doc, m, b)

    # ---------------- Ek-A sözlük
    doc.add_heading("Ek-A. Terimler sözlüğü", 1)
    _tablo(doc, ["Terim", "Anlamı"], [
        ["Kompartıman", "Tavan ve taban dahil her yanı en az 60 dk yangına dayanıklı elemanlarla ayrılmış bölge (Madde 4(tt))."],
        ["Kullanıcı yükü", "Bir anda binada veya bölümünde bulunma ihtimali olan toplam kişi sayısı."],
        ["Kullanıcı yük katsayısı", "Kişi başına düşen alan (m²/kişi), Ek-5/A."],
        ["Temiz genişlik", "Kapı/merdivende, kanat veya küpeşte çıkıntısı çıktıktan sonra kullanılabilir gerçek geçiş genişliği."],
        ["Kaçış uzaklığı", "Mekânın en uzak noktasından en yakın güvenli çıkışa kaçış yolu boyunca ölçülen mesafe."],
        ["Yapı yüksekliği", "Bodrum, asma kat ve çatı arası dahil tüm katların toplam yüksekliği."],
        ["Bina yüksekliği", "Binanın kot aldığı noktadan saçak seviyesine kadar olan mesafe."],
        ["Yüksek bina", "Bina yüksekliği >21,50 m veya yapı yüksekliği >30,50 m olan bina."],
        ["Tehlike sınıfı (DT/OT/YT)", "Yangın yükü ve yanabilirliğe göre Düşük, Orta (1-4), Yüksek (1-4) tehlike; sprinkler ve su hesabını belirler."],
        ["Islak / kuru sprinkler", "Boruları su dolu / basınçlı hava dolu sistem."],
        ["Hidrant", "Bina çevresindeki itfaiyenin bağlandığı yer üstü yangın musluğu."],
        ["mSS", "Metre su sütunu; ≈ 10 kPa."],
        ["Madde 92(5) usulü", "Su deposu = (sprinkler + dolap + hidrant debisi) × süre."],
        ["Hy", "Yangın pompasının sağlaması gereken basma yüksekliği."],
        ["MSDS / SDS", "Maddenin güvenlik bilgi formu (parlama noktası, sınıfı, depolama koşulları)."],
        ["Parlama noktası", "Sıvının üzerindeki buharın tutuşabildiği en düşük sıcaklık; yanıcı sıvı sınıfını belirler."],
        ["Havuzlama", "Tank sızıntısını tutan set/havuz hacmi."],
        ["REI / EI / E", "Yangına dayanım sembolleri: taşıma (R), bütünlük (E), yalıtım (I), Ek-3/A."],
    ], [4.5, 12.5], boyut=8.5)

    # ---------------- Ek-B madde eşleme
    doc.add_heading("Ek-B. Yönetmelik maddesi — program çıktısı eşlemesi", 1)
    _tablo(doc, ["Madde / Ek", "Konu", "Programda nerede"], [
        ["Madde 4, 7, 8, 14, 17-19, Ek-1", "Tanımlar, sınıflandırma", "Sonuçlar: 1. Sınıflandırma"],
        ["Madde 21-22", "Yerleşim, orman, itfaiye erişimi", "13. Çevre ve itfaiye erişimi"],
        ["Madde 23-28, Ek-3/B, 3/C", "Yangın dayanımı, cephe, çatı", "2. Yangın dayanımı"],
        ["Madde 24, Ek-4", "Kompartıman", "3. Yangın kompartımanı"],
        ["Madde 31-37, 39, 52, Ek-5/A-B, Ek-14", "Kaçış yolları ve mesafeleri", "4. Kullanıcı yükü ve kaçış yolları"],
        ["Madde 38-47", "Kaçış merdivenleri, kapılar", "5. Kaçış merdivenleri ve kapılar"],
        ["Madde 96, Ek-8/B", "Yağmurlama", "6. Yağmurlama sistemi"],
        ["Madde 91-93, Ek-8", "Su deposu, pompa", "7. Yangın suyu deposu ve pompa"],
        ["Madde 94-95, 97, 99", "Dolap, hidrant, söndürücü", "8. Yangın dolabı, hidrant ve söndürücüler"],
        ["Madde 70-77, 81, Ek-7", "Algılama, uyarı, aydınlatma", "9. Algılama, uyarı ve acil aydınlatma"],
        ["Madde 85-89", "Duman kontrolü, basınçlandırma", "10. Duman kontrolü ve basınçlandırma"],
        ["Madde 54-56, 65-66", "Kazan, yakıt, trafo, jeneratör", "11. Kazan dairesi, yakıt, trafo ve jeneratör"],
        ["Madde 101-123, Ek-9…12", "LPG, yanıcı sıvı", "12. Tehlikeli maddeler"],
        ["Madde 124-129", "Ekipler, eğitim", "14. Acil durum ekipleri"],
    ], [5.3, 5.2, 6.5], boyut=8.5)

    # ---------------- Ek-C kontrol listesi
    doc.add_heading("Ek-C. Rapor teslim öncesi kontrol listesi", 1)
    for m in ["Proje adı, firma, adres, tarih ve 'Hazırlayan' alanları dolduruldu.",
              "Kullanım sınıfı ve tehlike sınıfı kontrol edildi; birden fazla faaliyet varsa en yüksek sınıf seçildi.",
              "Kat adları Mahaller ve Kat kaçış verileri tablolarında aynı.",
              "Aynı anda kullanılmayan mahaller (tuvalet, soyunma, depo) 'Sayılır' işaretinden çıkarıldı (uygunsa).",
              "'Veri girilmedi' satırları işverenden tamamlatıldı veya raporda gerekçeyle bırakıldı.",
              "Varsayımlar (hidrant bölge riski, başlık alanı, eş zamanlı dolap) mühendisçe onaylandı.",
              "'Uygun değil' satırları için düzeltme önerileri hazırlandı ve işverene bildirildi.",
              "Ek-7, Ek-4 dipnotu gibi koşullu sonuçlar için itfaiye görüşü gerekliliği değerlendirildi.",
              "Excel ve Word raporu kaydedildi; proje .json dosyası arşivlendi.",
              "Rapor yönetmeliğin güncel hâline göre alındı (tablo sürümü kontrol edildi)."]:
        p = doc.add_paragraph("☐  " + m)
        p.paragraph_format.space_after = Pt(2)

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()
