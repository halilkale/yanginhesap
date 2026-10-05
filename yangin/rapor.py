"""Excel (.xlsx) ve Word (.docx) rapor üreticileri."""
from __future__ import annotations

import datetime as _dt
import io
from collections import OrderedDict

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import hesaplar as H
from . import tablolar as T
from .modeller import (
    BILGI, GEREKLI, GEREKMEZ, KONTROL, KOSULLU, UYGUN, UYGUN_DEGIL, Girdi, Sonuc,
)

YONETMELIK = "Binaların Yangından Korunması Hakkında Yönetmelik (RG 19.12.2007/26735; 20.11.2021 tarihli 4825 sayılı CK ile değişik)"
KILAVUZ = "Binaların Yangından Korunması Hakkında Yönetmelik Kılavuzu (ÇŞİDB, Aralık 2024)"
UYARI = (
    "Bu rapor, yönetmelik maddeleri ve eklerindeki sayısal kriterlere göre yapılan ön hesap ve kontrolleri içerir. "
    "Projenin onayı, itfaiye teşkilatının görüşü, ilgili Türk Standartları (TS EN 12845, TS EN 54, TS EN 671, TS EN 81-72 vb.) "
    "ve yetkili meslek mensuplarının sorumluluğunda yürütülen detay tasarımın yerine geçmez. "
    "Kılavuzda belirtildiği üzere mevzuatla çelişme halinde yürürlükteki mevzuat geçerlidir."
)

DURUM_RENK = {
    UYGUN: "C6EFCE",
    UYGUN_DEGIL: "FFC7CE",
    GEREKLI: "DDEBF7",
    KOSULLU: "FFEB9C",
    KONTROL: "EDEDED",
    GEREKMEZ: "F2F2F2",
    BILGI: "FFFFFF",
}


def _ana_bolum(b: str) -> str:
    return b.split(" — ")[0]


def _gruplar(sonuclar: list[Sonuc]) -> "OrderedDict[str, list[Sonuc]]":
    d: "OrderedDict[str, list[Sonuc]]" = OrderedDict()
    for r in sonuclar:
        d.setdefault(_ana_bolum(r.bolum), []).append(r)
    return d


def _fmt(v) -> str:
    if isinstance(v, float):
        if abs(v - round(v)) < 1e-9:
            return f"{int(round(v)):,}".replace(",", ".")
        return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if isinstance(v, int):
        return f"{v:,}".replace(",", ".")
    return str(v)


def girdi_ozeti(g: Girdi) -> list[tuple[str, str]]:
    ev = lambda b: "Evet" if b else "Hayır"
    return [
        ("Proje adı", g.proje_adi),
        ("Tesis / firma", g.tesis_adi),
        ("Adres", g.adres),
        ("Hazırlayan", g.hazirlayan),
        ("Tarih", g.tarih or _dt.date.today().strftime("%d.%m.%Y")),
        ("Kullanım sınıfı", T.KULLANIM_SINIFLARI[g.kullanim]),
        ("Tehlike sınıfı", T.TEHLIKE_SINIFLARI[g.tehlike]),
        ("Yapı türü", "Mevcut yapı" if g.mevcut_yapi else "Yeni yapı"),
        ("Taşıyıcı sistem", {"betonarme": "Betonarme", "celik": "Çelik", "ahsap": "Ahşap", "kagir": "Kâgir"}[g.tasiyici]),
        ("Yapı yüksekliği", f"{g.yapi_yuksekligi:g} m"),
        ("Bina yüksekliği", f"{g.bina_yuksekligi:g} m"),
        ("Kat sayısı (zemin üstü / bodrum)", f"{g.kat_sayisi} / {g.bodrum_kat_sayisi}"),
        ("Taban alanı", f"{g.taban_alani:,.0f} m²"),
        ("Toplam kapalı alan", f"{g.toplam_kapali_alan:,.0f} m²"),
        ("En büyük kat / kompartıman alanı", f"{g.en_buyuk_kat_alani:,.0f} m² / {g.ozet_alan():,.0f} m²"),
        ("Bina boyutları (boy × en)", f"{g.bina_boyu:g} × {g.bina_eni:g} m"),
        ("Kolay alevlenici/parlayıcı madde", ev(g.kolay_alevlenici)),
        ("Yağmurlama sistemi", f"{'Var (' + ('ıslak' if g.yagmurlama_tipi == 'islak' else 'kuru') + ')' if g.yagmurlama_var else 'Yok'}"),
        ("Otomatik algılama", ev(g.algilama_var)),
        ("Duman tahliye sistemi", ev(g.duman_tahliye_var)),
        ("Hidrant / yangın dolabı", f"{ev(g.hidrant_var)} / {ev(g.dolap_var)}"),
        ("Orman alanı içinde/bitişiğinde", ev(g.orman_yakin)),
    ]


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def excel_olustur(g: Girdi, sonuclar: list[Sonuc] | None = None) -> bytes:
    sonuclar = sonuclar if sonuclar is not None else H.hesapla(g)
    wb = Workbook()
    ince = Side(style="thin", color="BFBFBF")
    kenar = Border(left=ince, right=ince, top=ince, bottom=ince)
    baslik_fill = PatternFill("solid", fgColor="1F3864")
    baslik_font = Font(bold=True, color="FFFFFF")

    # --- Özet
    ws = wb.active
    ws.title = "Özet"
    ws["A1"] = "YANGIN KORUMA HESAP RAPORU"
    ws["A1"].font = Font(bold=True, size=16, color="1F3864")
    ws["A2"] = YONETMELIK
    ws["A2"].font = Font(italic=True, size=9)
    r = 4
    for k, v in girdi_ozeti(g):
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v)
        r += 1
    r += 1
    ws.cell(r, 1, "SONUÇ ÖZETİ").font = Font(bold=True, size=12, color="1F3864")
    r += 1
    oz = H.ozet(sonuclar)
    for durum in (UYGUN, UYGUN_DEGIL, GEREKLI, KOSULLU, KONTROL, GEREKMEZ, BILGI):
        if durum in oz:
            c = ws.cell(r, 1, durum)
            c.fill = PatternFill("solid", fgColor=DURUM_RENK[durum])
            c.font = Font(bold=True)
            ws.cell(r, 2, oz[durum])
            r += 1
    r += 1
    ws.cell(r, 1, "UYARI").font = Font(bold=True, color="C00000")
    ws.cell(r, 2, UYARI).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 75
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 90

    # --- Sonuçlar
    ws2 = wb.create_sheet("Sonuçlar")
    basliklar = ["Bölüm", "Kalem", "Değer", "Birim", "Durum", "Madde / Ek", "Açıklama / Formül"]
    for j, b in enumerate(basliklar, 1):
        c = ws2.cell(1, j, b)
        c.fill, c.font, c.border = baslik_fill, baslik_font, kenar
        c.alignment = Alignment(horizontal="center", vertical="center")
    for i, x in enumerate(sonuclar, 2):
        vals = [x.bolum, x.kalem, x.deger, x.birim, x.durum, x.madde, x.aciklama]
        for j, v in enumerate(vals, 1):
            c = ws2.cell(i, j, v)
            c.border = kenar
            c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right" if (j == 3 and isinstance(v, (int, float))) else "left")
        ws2.cell(i, 5).fill = PatternFill("solid", fgColor=DURUM_RENK.get(x.durum, "FFFFFF"))
        ws2.cell(i, 5).font = Font(bold=x.durum == UYGUN_DEGIL)
        if isinstance(x.deger, float):
            ws2.cell(i, 3).number_format = "#,##0.00"
    for j, w in enumerate([34, 52, 18, 8, 16, 24, 90], 1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    ws2.freeze_panes = "C2"
    ws2.auto_filter.ref = f"A1:G{len(sonuclar) + 1}"

    # --- Uygun değil / eksik
    ws3 = wb.create_sheet("Eksikler")
    ws3.append(["Bölüm", "Kalem", "Değer", "Birim", "Durum", "Madde / Ek", "Açıklama"])
    for j in range(1, 8):
        c = ws3.cell(1, j)
        c.fill, c.font = baslik_fill, baslik_font
    for x in sonuclar:
        if x.durum in (UYGUN_DEGIL, KONTROL):
            ws3.append([x.bolum, x.kalem, x.deger, x.birim, x.durum, x.madde, x.aciklama])
            ws3.cell(ws3.max_row, 5).fill = PatternFill("solid", fgColor=DURUM_RENK[x.durum])
            for j in range(1, 8):
                ws3.cell(ws3.max_row, j).alignment = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate([34, 52, 18, 8, 16, 24, 90], 1):
        ws3.column_dimensions[get_column_letter(j)].width = w
    ws3.freeze_panes = "A2"

    # --- Kullanıcı yükü dökümü
    if g.katlar:
        ws4 = wb.create_sheet("Kullanıcı yükü")
        ws4.append(["Kat", "Mahal", "Alan (m²)", "Katsayı (m²/kişi)", "Belirli kişi sayısı", "Hesapta sayılır", "Kişi"])
        for j in range(1, 8):
            c = ws4.cell(1, j)
            c.fill, c.font = baslik_fill, baslik_font
        for kat in g.katlar:
            for m in kat.mahaller:
                n = (m.alan / m.katsayi) if m.katsayi > 0 else 0
                if m.kisi_belirli:
                    n = max(n, m.kisi_belirli)
                ws4.append([kat.ad, m.ad, m.alan, m.katsayi, m.kisi_belirli or "", "Evet" if m.sayilir else "Hayır", round(n, 2) if m.sayilir else 0])
            ws4.append([kat.ad, "TOPLAM", "", "", "", "", round(H.kat_kullanici_yuku(kat), 2)])
            for j in range(1, 8):
                ws4.cell(ws4.max_row, j).font = Font(bold=True)
        for j, w in enumerate([22, 40, 14, 18, 18, 16, 12], 1):
            ws4.column_dimensions[get_column_letter(j)].width = w

    # --- Dayanak
    ws5 = wb.create_sheet("Dayanak")
    ws5.append(["Kaynak"])
    ws5.append([YONETMELIK])
    ws5.append([KILAVUZ])
    ws5.append([""])
    ws5.append(["Danıştay 10. Daire kararı (E.2019/261, K.2021/5537; İDDK 22/2/2023 E.2022/2463, K.2023/317) ile Ek-7'nin endüstriyel, ticaret, kurum ve toplanma amaçlı yapılara ilişkin kısımları iptal edilmiştir."])
    ws5.append([""])
    ws5.append([UYARI])
    ws5.column_dimensions["A"].width = 140
    for row in ws5.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------
def _golge(cell, hex_renk: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_renk)
    for eski in tcPr.findall(qn("w:shd")):
        tcPr.remove(eski)
    # CT_TcPr şema sırası: ... tcBorders, shd, noWrap, tcMar, textDirection, tcFitText, vAlign
    tcPr.insert_element_before(shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")


def _hucre(cell, metin, kalin=False, boyut=8.5, renk=None, hizala=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(metin))
    run.font.size = Pt(boyut)
    run.bold = kalin
    if renk:
        run.font.color.rgb = RGBColor.from_string(renk)
    if hizala:
        p.alignment = hizala


def _tablo_baslik_tekrar(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def word_olustur(g: Girdi, sonuclar: list[Sonuc] | None = None) -> bytes:
    sonuclar = sonuclar if sonuclar is not None else H.hesapla(g)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    stil = doc.styles["Normal"]
    stil.font.name = "Calibri"
    stil.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    stil.font.size = Pt(10)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("YANGIN KORUMA HESAP RAPORU")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    t2 = doc.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t2.add_run(YONETMELIK)
    r.italic = True
    r.font.size = Pt(9)

    doc.add_heading("1. Proje bilgileri", level=1)
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = "Table Grid"
    for k, v in girdi_ozeti(g):
        row = tbl.add_row().cells
        _hucre(row[0], k, kalin=True, boyut=9)
        _hucre(row[1], v, boyut=9)
        _golge(row[0], "F2F2F2")
        row[0].width, row[1].width = Cm(5.5), Cm(11.8)

    doc.add_heading("2. Sonuç özeti", level=1)
    oz = H.ozet(sonuclar)
    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = "Table Grid"
    _hucre(tbl.rows[0].cells[0], "Durum", kalin=True)
    _hucre(tbl.rows[0].cells[1], "Adet", kalin=True)
    for durum in (UYGUN, UYGUN_DEGIL, GEREKLI, KOSULLU, KONTROL, GEREKMEZ, BILGI):
        if durum in oz:
            c = tbl.add_row().cells
            _hucre(c[0], durum, kalin=True)
            _hucre(c[1], oz[durum])
            _golge(c[0], DURUM_RENK[durum])

    kritik = [x for x in sonuclar if x.durum == UYGUN_DEGIL]
    eksik = [x for x in sonuclar if x.durum == KONTROL]
    doc.add_heading("3. Uygun olmayan hususlar", level=1)
    if kritik:
        for x in kritik:
            p = doc.add_paragraph(style="List Bullet")
            a = p.add_run(f"{x.kalem}: ")
            a.bold = True
            p.add_run(f"{_fmt(x.deger)} {x.birim} — {x.madde}. {x.aciklama}")
    else:
        doc.add_paragraph("Girilen verilere göre uygun olmayan husus tespit edilmemiştir.")
    if eksik:
        doc.add_heading("Girilmeyen / doğrulanması gereken veriler", level=2)
        for x in eksik:
            p = doc.add_paragraph(style="List Bullet")
            a = p.add_run(f"{x.kalem}: ")
            a.bold = True
            p.add_run(f"{x.madde}. {x.aciklama}")

    doc.add_heading("4. Hesap ve kontroller", level=1)
    for n, (bolum, liste) in enumerate(_gruplar(sonuclar).items(), 1):
        doc.add_heading(bolum, level=2)
        tbl = doc.add_table(rows=1, cols=4)
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        genislik = [Cm(5.0), Cm(3.4), Cm(2.6), Cm(6.3)]
        for j, b in enumerate(["Kalem", "Değer", "Durum", "Madde / Açıklama"]):
            _hucre(tbl.rows[0].cells[j], b, kalin=True, renk="FFFFFF")
            _golge(tbl.rows[0].cells[j], "1F3864")
        _tablo_baslik_tekrar(tbl.rows[0])
        onceki = None
        for x in liste:
            if x.bolum != bolum and x.bolum != onceki:
                satir = tbl.add_row()
                m = satir.cells[0].merge(satir.cells[3])
                _hucre(m, x.bolum.split(" — ", 1)[1], kalin=True, boyut=9)
                _golge(m, "D9E1F2")
                onceki = x.bolum
            c = tbl.add_row().cells
            _hucre(c[0], x.kalem)
            _hucre(c[1], (_fmt(x.deger) + (f" {x.birim}" if x.birim else "")).strip(), kalin=True)
            _hucre(c[2], x.durum, kalin=x.durum == UYGUN_DEGIL, boyut=8)
            _golge(c[2], DURUM_RENK.get(x.durum, "FFFFFF"))
            metin = x.madde + (f" — {x.aciklama}" if x.aciklama else "")
            _hucre(c[3], metin, boyut=8)
            for j, w in enumerate(genislik):
                c[j].width = w

    doc.add_heading("5. Dayanak ve uyarılar", level=1)
    doc.add_paragraph(f"Dayanak: {YONETMELIK}.")
    doc.add_paragraph(f"Açıklayıcı kaynak: {KILAVUZ}.")
    doc.add_paragraph("Danıştay 10. Daire kararı (E.2019/261, K.2021/5537; İDDK 22/2/2023, E.2022/2463, K.2023/317) ile Ek-7 tablosunun endüstriyel, ticaret, kurum ve toplanma amaçlı yapılara ilişkin kısımları iptal edilmiştir; rapor bu hali esas almıştır.")
    p = doc.add_paragraph()
    a = p.add_run("Uyarı: ")
    a.bold = True
    a.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    p.add_run(UYARI)

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()
