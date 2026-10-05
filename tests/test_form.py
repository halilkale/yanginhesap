"""İşveren formu: şablon üretimi, doldurma ve programa aktarma gidiş-dönüş testleri."""
import io

import pytest
from openpyxl import load_workbook

from yangin import alanlar as A
from yangin import form as F
from yangin import hesaplar as H
from yangin.modeller import Girdi
from yangin.ornek import ornek_fabrika


def test_form_tum_alanlari_kapsar():
    assert A.kapsam_kontrol() == []


def test_sablon_olusur_ve_acilir():
    wb = load_workbook(io.BytesIO(F.sablon_olustur()))
    for s in A.SAYFALAR:
        assert s.ad in wb.sheetnames
    for t in A.TABLO_SAYFALARI:
        assert t.ad in wb.sheetnames
    assert "Belgeler" in wb.sheetnames and "Talimat" in wb.sheetnames


def test_bos_sablon_aktarilir_ve_eksikleri_bildirir():
    g, uyari, eksik = F.formdan_oku(F.sablon_olustur())
    assert isinstance(g, Girdi)
    assert any("Mahal" in e or "mahal" in e for e in eksik)
    assert len(eksik) > 5
    assert uyari == []        # örnek satırlar atlanmalı


def test_gidis_donus_ayni_sonuc():
    g0 = ornek_fabrika()
    g1, uyari, eksik = F.formdan_oku(F.doldurulmus_form(g0))
    assert uyari == []
    a = H.hesapla(g0)
    b = H.hesapla(g1)
    assert [(x.kalem, x.deger, x.durum) for x in a] == [(x.kalem, x.deger, x.durum) for x in b]


def test_tesis_secimi_tehlikeyi_belirler():
    g0 = ornek_fabrika()
    veri = F.doldurulmus_form(g0)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(veri))
    ws = wb["1_Genel"]
    for r in range(5, 30):
        if ws.cell(r, F.ANAHTAR_SUTUN).value == "tehlike":
            ws.cell(r, 5).value = None
        if ws.cell(r, F.ANAHTAR_SUTUN).value == "@tesis":
            ws.cell(r, 5).value = A.TESIS_ETIKET[-1]       # Havai fişek fabrikası → YT4
    bio = io.BytesIO(); wb.save(bio)
    g, _, _ = F.formdan_oku(bio.getvalue())
    assert g.tehlike == "YT4"


def test_hatali_deger_uyari_verir():
    wb = load_workbook(io.BytesIO(F.doldurulmus_form(ornek_fabrika())))
    ws = wb["2_Bina_Olculeri"]
    ws.cell(5, 5).value = "yüz metre"          # yapı yüksekliği
    bio = io.BytesIO(); wb.save(bio)
    g, uyari, _ = F.formdan_oku(bio.getvalue())
    assert any("anlaşılamadı" in u for u in uyari)


def test_virgullu_sayi():
    assert F._sayi("12,5") == 12.5 and F._sayi("1.250,5") == 1250.5 and F._sayi(" ") is None
