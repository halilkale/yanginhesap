"""Hesap dökümü ile hesap motorunun tutarlılığı."""
from yangin import dokum as D
from yangin import hesaplar as H
from yangin.modeller import Girdi, Kapi, Kat, Mahal
from yangin.ornek import ornek_fabrika


def bloklar(g):
    return {b.baslik: b for bo in D.hesap_dokumu(g) for b in bo.bloklar}


def test_kilavuz_ornegi_dokumde():
    g = Girdi(kullanim="buro", tehlike="OT2", yagmurlama_var=True, es_zamanli_dolap=2, hidrant_var=True, dolap_var=True,
              statik_yukseklik_mSS=20, boru_kaybi_mSS=6.06, akma_basinci_mSS=70)
    g.katlar = [Kat(ad="Tipik", mahaller=[Mahal("Ofis", 2026.4, 10)], cikis_turu="merdiven")]
    b = bloklar(g)
    assert "79,2 m³" in b["Su deposu hacmi"].sonuc or "105" in b["Su deposu hacmi"].sonuc
    assert any("1.320 l/dk" in a for a in b["Su deposu hacmi"].adimlar)
    assert any("168,87 cm" in a for a in b["Çıkış genişliği — Tipik"].adimlar)
    assert any("101,32 cm" in a for a in b["Çıkış genişliği — Tipik"].adimlar)
    assert any("96,06 mSS" in a for a in b["Yangın pompası"].adimlar)


def test_kullanici_yuku_dokumu():
    g = ornek_fabrika()
    b = bloklar(g)["Kullanıcı yükü — Zemin kat"]
    assert b.sonuc == "610 kişi"
    assert any("4.800 m² ÷ 10 m²/kişi = 480 kişi" in a for a in b.adimlar)


def test_durumlar_hesap_motoruyla_ayni():
    g = ornek_fabrika()
    S = H.hesapla(g)
    b = bloklar(g)
    ks = next(r for r in S if r.kalem.startswith("Ölçülen en uzak") and "Zemin" in r.bolum)
    assert b["Kaçış uzaklığı — Zemin kat"].durum == ks.durum == "UYGUN DEĞİL"
    kom = next(r for r in S if r.kalem == "Kompartıman alanı kontrolü")
    assert b["Kompartıman alanı"].durum == kom.durum
    su = next(r for r in S if r.kalem == "Önerilen en az yangın suyu deposu")
    assert D.n(su.deger, 1) in b["Su deposu hacmi"].sonuc


def test_kapi_tek_kanatli_160_aciklamasi():
    g = Girdi()
    g.kapilar = [Kapi("K1", 160, 210, 1), Kapi("K2", 160, 210, 2), Kapi("K3", 160, 210, 0)]
    S = {r.bolum.split("Kapı ")[-1]: r for r in H.merdiven_kapi(g) if r.kalem == "Temiz genişlik"}
    assert S["K1"].durum == "UYGUN DEĞİL" and "iki kanatlı" in S["K1"].aciklama and "120 cm" in S["K1"].aciklama
    assert S["K2"].durum == "UYGUN"
    assert S["K3"].durum == "UYGUN DEĞİL" and "girilmedi" in S["K3"].aciklama      # kanat girilmediyse tek kanatlı kabul


def test_dokum_duz_metin_ve_raporlar():
    from yangin import rapor
    import io
    from openpyxl import load_workbook
    g = ornek_fabrika()
    t = D.dokum_duz_metin(D.hesap_dokumu(g))
    assert "SONUÇ" in t and "610 kişi" in t
    wb = load_workbook(io.BytesIO(rapor.excel_olustur(g)))
    assert "Hesap Dökümü" in wb.sheetnames and wb["Hesap Dökümü"].max_row > 50
    assert len(rapor.word_olustur(g)) > 10000
