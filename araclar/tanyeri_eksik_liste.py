"""Tanyeri (Tuzla) için hesabın tamamlanması gereken verilerin kontrol listesi (Excel)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# (öncelik, konu, istenen bilgi, neden gerekli / hangi hesabı değiştirir, kimden / hangi belge, şimdiki durum)
S = [
 (1, "Kat alanları", "3. katın gerçek alanı; 2. katın (2.240,73 m²) teyidi; zemin ve 1. katın teyidi", "Toplam kapalı alan, kompartıman ve söndürücü/dolap/sprinkler adetleri, Madde 7(4) (10.000 m²) eşiği", "Mimar / ruhsat alan tablosu", "3. kat TAHMİN (2.240,73); toplam 7.621,64 m²"),
 (1, "Kompartıman planı", "Her kompartımanın sınırları ve alanı (özellikle Depo üst: tek 1.403 m² mi, bölünmüş mü?); kompartıman duvar/döşeme ve kapı yangın dayanımları (REI / E süreleri)", "Ek-4 alan sınırı (depo OT-3: 1.000 m²), Madde 24, Ek-3/B", "Yangın projesi / kompartıman çalışması", "Mahal numarası eşleştirmesi varsayım; dayanımlar bilinmiyor"),
 (1, "Aktif yangın sistemleri", "Sprinkler var/yok (tip, tasarım yoğunluğu, koruma alanı, en alt–en üst başlık yüksekliği h); yangın algılama ve alarm; duman tahliye (doğal/mekanik)", "Madde 96(2)e (kolay alevlenici + >1.000 m² → sprinkler zorunlu); Ek-4 dipnotları (kompartıman sınırsız olur mu); Ek-3/C, Ek-5/B", "Mekanik/elektrik projeleri, işveren", "Sprinkler/algılama/duman: YOK varsayıldı"),
 (1, "Stok ve malzeme", "Azami (anlık değil) stok miktarları kg; plastiğin cinsi (PP/PE/ABS…); ambalaj payı; boyanın MSDS ısıl değeri (MJ/kg) ve parlama noktası; boyahane/depodaki solvent ve tiner", "Yangın yükü (MJ/m²), yanıcı sıvı sınıfı/eşikleri (Ek-11, Madde 114–118), kolay alevlenici madde kararı", "Üretim/depo müdürü, MSDS", "Verilen kg'lar küçük; 43 MJ/kg ve 25 MJ/kg varsayıldı"),
 (1, "Kaçış verileri", "Her katta çıkış sayısı ve türü, çıkışların toplam ve en dar temiz genişliği, en uzak noktadan çıkışa kaçış mesafesi (tek/iki yön), çıkmaz koridor, çıkışlar arası mesafe", "Madde 32–33, 39, Ek-5/B — şu an tüm katlar için kontrol YAPILAMIYOR", "Yangın tahliye planı + saha ölçümü", "Hiç girilmedi (VERİ GİRİLMEDİ)"),
 (1, "Kaçış merdivenleri ve kapılar", "Her merdiven: temiz genişlik, rıht, basış, sahanlık basamak sayısı, baş yüksekliği, kapı ve duvar yangın dayanımı, hizmet verdiği kat sayısı; başlıca çıkış kapıları (genişlik, kanat sayısı, açılış yönü, panik bar)", "Madde 33, 38, 41, 47", "Mimari proje, saha ölçümü, kapı etiketleri", "Hiç girilmedi"),
 (1, "Yükseklik ve kotlar", "Bina ve yapı yüksekliği (saçak kotu, itfaiye yaklaşma kotu: planda +0,00 ve +4,50 var); bodrum kat var mı, derinliği", "Yüksek bina testi, Ek-3/C sütunu, basınçlandırma, itfaiye su verme", "Mimari kesit/cephe", "13 m TAHMİN, bodrum yok varsayıldı"),
 (1, "Çalışan sayısı", "Vardiya başına en yüksek çalışan sayısı; hat başı gibi kişi sayısı belli mahaller (enjeksiyon, montaj, boyahane)", "Kullanıcı yükü, çıkış genişliği, Madde 126 (acil durum ekipleri)", "İnsan kaynakları", "Girilmedi (mahal alanından hesaplanıyor: 1. kat 139, 3. kat 148 kişi vb.)"),
 (2, "Taşıyıcı sistem", "Betonarme/çelik; betonarmede kolon-kiriş-döşeme net beton paspayı; çelikse yangın koruması; cephe ve çatı kaplama malzemesi (yanıcılık sınıfı)", "Madde 23–28, Ek-3/B-C yangın dayanım süreleri", "Statik proje, cephe/çatı detayları", "Betonarme varsayıldı; paspayı yok"),
 (2, "Yangın suyu ve pompa", "Yangın suyu deposu hacmi; pompa etiketi ve karakteristik eğrisi (anma debi/basma, kapalı vana, %150 debi); asıl/yedek pompa adedi; statik yükseklik ve kritik hat boru kaybı; yangın dolabı adedi/tipi", "Madde 92–94 su deposu ve pompa kontrolleri (Hy)", "Mekanik proje, pompa etiketi", "Planda 'YANGIN POMPASI' yazıyor; değerler yok"),
 (2, "Hidrant ve itfaiye erişimi", "Dış hidrant adedi ve yerleri, itfaiye su verme bağlantısı; itfaiye aracının son yaklaşma noktasından cepheye mesafe; iç yol genişliği, dönemeç yarıçapları, eğim, serbest yükseklik, taşıma yükü", "Madde 22, 95, 97", "Vaziyet planı, saha", "Girilmedi"),
 (2, "Acil aydınlatma ve uyarı", "Acil aydınlatma süresi ve lux ölçümü, yönlendirme levhası boyu/tipi, en uzun kaçış yolu uzunluğu, yangın uyarı butonu adedi", "Madde 72–75", "Elektrik projesi, ölçüm raporu", "Girilmedi"),
 (2, "Kazan dairesi ve enerji", "Z14 kazan dairesi (35,96 m²): ısıl kapasite (kW), yakıt türü, çıkış kapısı sayısı, duvar/kapı dayanımı; yakıt/LPG/doğalgaz; trafo (yağlı/kuru) ve jeneratör odası", "Madde 54–56, 65–66, 94", "Mekanik/elektrik proje", "Kazan yok varsayıldı"),
 (2, "Boya ve yanıcı sıvı depolama", "Sıvı boya/tiner nerede ve ne kadar saklanıyor; kap/tank cinsi; depo odası alanı, dayanımı, havalandırma ve yangın korunumu; boyahane elektrik (ex-proof) ve havalandırma", "Madde 113–118, 88(5), Ek-11, Ek-12", "Boyahane sorumlusu, MSDS", "Girilmedi"),
 (2, "Mahal kullanımları", "'Rezerv alan' mahallerinin (212: 354 m², 310: 1.403 m²) gerçek kullanımı ve kişi yoğunluğu; tuvalet/soyunma gibi aynı anda kullanılmayan mahaller", "Kullanıcı yükü (Ek-5/A katsayısı)", "İşveren", "Depo (30 m²/kişi) varsayıldı"),
 (3, "Yapı durumu", "Yeni yapı mı, mevcut işletme mi (ruhsat/iskân tarihi, imar durumu); parseldeki diğer binalar", "Onuncu Kısım (mevcut yapı), Madde 95(7) (toplam taban >5.000 m² → dış hidrant)", "Ruhsat, iskân", "Yeni yapı, tek bina varsayıldı"),
 (3, "Çevre", "Orman sınırına yakınlık ve arazi eğimi", "Madde 7(12), 21(5)", "İmar durumu", "Orman alanı değil varsayıldı"),
 (3, "Organizasyon", "Yangın güvenliği sorumlusu, acil durum ekipleri (söndürme/kurtarma/koruma/ilk yardım), tatbikat kayıtları", "Madde 124–129", "İSG", "Bilinmiyor"),
 (3, "Belgeler", "Yangın tahliye planı; önceki itfaiye raporu/tutanağı; asansör bilgisi", "Genel kontrol", "İşveren/belediye", "—"),
]
hedef = os.path.join(os.path.dirname(__file__), "..", "belgeler", "tanyeri", "Tanyeri_Eksik_Veri_Listesi.xlsx")
wb = Workbook(); ws = wb.active; ws.title = "Eksik veriler"
ws["A1"] = "TANYERİ (TUZLA) — HESABIN TAMAMLANMASI İÇİN GEREKEN VERİLER"; ws["A1"].font = Font(bold=True, size=14, color="1F3864")
ws["A2"] = "Öncelik 1: sonucu değiştirir/hesabı engeller.  Öncelik 2: kontrolün yapılabilmesi için gerekli.  Öncelik 3: tamamlayıcı. Sarı sütunları işletme doldurur."
ince = Side(style="thin", color="A6A6A6"); kenar = Border(left=ince, right=ince, top=ince, bottom=ince)
baslik = ["Öncelik", "Konu", "İstenen bilgi", "Neden gerekli / hangi hesabı etkiler", "Kimden / hangi belge", "Şimdiki durum (varsayım)", "CEVAP / VERİ", "Durum"]
for j, b in enumerate(baslik, 1):
    c = ws.cell(4, j, b); c.fill = PatternFill("solid", fgColor="1F3864"); c.font = Font(bold=True, color="FFFFFF"); c.border = kenar
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
renk = {1: "F8CBAD", 2: "FFE699", 3: "DDEBF7"}
dv = DataValidation(type="list", formula1='"Teslim edildi,Bekleniyor,Bu tesiste yok"', allow_blank=True); ws.add_data_validation(dv)
for i, (p, konu, ist, neden, kim, simdi) in enumerate(S, 5):
    for j, v in enumerate([p, konu, ist, neden, kim, simdi, None, None], 1):
        c = ws.cell(i, j, v); c.border = kenar; c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(i, 1).fill = PatternFill("solid", fgColor=renk[p]); ws.cell(i, 1).alignment = Alignment(horizontal="center", vertical="top")
    ws.cell(i, 7).fill = PatternFill("solid", fgColor="FFF2CC"); ws.cell(i, 8).fill = PatternFill("solid", fgColor="FFF2CC"); dv.add(f"H{i}")
    ws.row_dimensions[i].height = max(48, 15 * (len(ist) // 55 + 1))
for j, w in enumerate([9, 24, 62, 48, 28, 34, 40, 18], 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = "C5"; ws.auto_filter.ref = f"A4:H{4 + len(S)}"
wb.save(hedef); print("yazıldı", len(S), "kalem")
